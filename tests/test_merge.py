"""End-to-end tests: build workbook versions with XlsxWriter (which writes
Excel-style files, charts and all), commit them on branches, and let real
`git merge` run xlgit as the merge driver.

Set XLGIT_KEEP=<dir> to keep every merged workbook for opening in Excel.
"""
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
import xlsxwriter
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import xlgit  # noqa: E402

XLGIT = Path(xlgit.__file__).resolve()


def build(path, *, b2=1000, c3=350, title="Spending", budget_chart=True, notes_name="Notes",
          notes_chart=False, notes_title="Q1 only", forecast=False, food_name=False,
          note_text="check with landlord"):
    wb = xlsxwriter.Workbook(str(path))
    bold = wb.add_format({"bold": True})
    money = wb.add_format({"num_format": "$#,##0"})
    ws = wb.add_worksheet("Budget")
    ws.set_column(0, 0, 18)
    ws.write_row(0, 0, ["Item", "Q1", "Q2", "Total"], bold)
    for r, (item, q1, q2) in enumerate([("Rent", b2, 1000), ("Food", 300, c3)], start=1):
        ws.write(r, 0, item)
        ws.write_number(r, 1, q1, money)
        ws.write_number(r, 2, q2, money)
        ws.write_formula(r, 3, f"=B{r + 1}+C{r + 1}", money)
    ws.write(3, 0, "Total", bold)
    for col in "BCD":
        ws.write_formula(f"{col}4", f"=SUM({col}2:{col}3)", money)
    if budget_chart:
        ch = wb.add_chart({"type": "column"})
        for col in "BC":
            ch.add_series({"name": f"=Budget!${col}$1", "categories": "=Budget!$A$2:$A$3",
                           "values": f"=Budget!${col}$2:${col}$3"})
        ch.set_title({"name": title})
        ws.insert_chart("F2", ch)
    notes = wb.add_worksheet(notes_name)
    notes.write("A1", "draft")
    notes.write_comment("A1", note_text)
    if notes_chart:
        ch = wb.add_chart({"type": "line"})
        ch.add_series({"categories": "=Budget!$A$2:$A$3", "values": "=Budget!$B$2:$B$3"})
        ch.set_title({"name": notes_title})
        notes.insert_chart("C2", ch)
    if forecast:
        f = wb.add_worksheet("Forecast")
        f.write_row(0, 0, ["Month", "Cash"], bold)
        for i in range(1, 7):
            f.write(i, 0, f"M{i}")
            f.write_number(i, 1, 1000 + 150 * i, money)
        ch = wb.add_chart({"type": "line"})
        ch.add_series({"categories": "=Forecast!$A$2:$A$7", "values": "=Forecast!$B$2:$B$7"})
        ch.set_title({"name": "Cash forecast"})
        f.insert_chart("D2", ch)
    wb.define_name("Rent", "=Budget!$B$2")
    if food_name:
        wb.define_name("Food", "=Budget!$C$3")
    wb.close()


class Repo:
    def __init__(self, root):
        self.root = root
        self.git("init", "-q", "-b", "main")
        self.git("config", "user.email", "t@t")
        self.git("config", "user.name", "t")
        subprocess.run([sys.executable, str(XLGIT), "install", "--repo"], cwd=root, check=True, capture_output=True)
        self.book = root / "budget.xlsx"

    def git(self, *args, check=True):
        return subprocess.run(["git", *args], cwd=self.root, check=check, capture_output=True, text=True)

    def commit(self, branch, from_branch, **variant):
        self.git("checkout", "-q", from_branch, check=branch != from_branch)  # no commits yet on first call
        if branch != from_branch:
            self.git("checkout", "-q", "-b", branch)
        build(self.book, **variant)
        self.git("add", "-A")
        self.git("commit", "-q", "-m", branch)

    def merge(self, ours, theirs):
        self.git("checkout", "-q", ours)
        return self.git("merge", "--no-edit", theirs, check=False)


BASE = {}
ALICE = dict(b2=1100, title="Spending 2026", notes_name="Notes 2026")
BOB = dict(c3=400, notes_chart=True, forecast=True, food_name=True)


@pytest.fixture
def repo(tmp_path):
    r = Repo(tmp_path)
    r.commit("main", "main", **BASE)
    r.commit("alice", "main", **ALICE)
    r.commit("bob", "main", **BOB)
    return r


def keep(path, name):
    out = os.environ.get("XLGIT_KEEP")
    if out:
        os.makedirs(out, exist_ok=True)
        shutil.copy(path, os.path.join(out, name))


def objects(path):
    return xlgit.describe_objects(xlgit.Package.open(str(path)))


def charts(path, sheet):
    return [o.split("'")[1] for o in objects(path).get(sheet, []) if o.startswith("chart:")]


def assert_opens_clean(path):
    """Structural checks: loads in openpyxl, every relationship target exists,
    every part has a content type."""
    load_workbook(path)
    pkg = xlgit.Package.open(str(path))
    for owner in [""] + list(pkg.parts):
        for _, _, tgt, ext in pkg.rels(owner):
            assert ext or tgt in pkg.parts, f"{owner} links to missing {tgt}"
    for name in pkg.parts:
        if name != "[Content_Types].xml":
            assert pkg.content_type(name), f"no content type for {name}"


@pytest.mark.parametrize("ours,theirs", [("bob", "alice"), ("alice", "bob")])
def test_clean_merge_keeps_every_chart(repo, ours, theirs):
    result = repo.merge(ours, theirs)
    assert result.returncode == 0, result.stdout + result.stderr
    keep(repo.book, f"merged_{ours}_{theirs}.xlsx")
    assert_opens_clean(repo.book)

    wb = load_workbook(repo.book)
    assert wb.sheetnames == ["Budget", "Notes 2026", "Forecast"]
    assert wb["Budget"]["B2"].value == 1100  # alice
    assert wb["Budget"]["C3"].value == 400   # bob
    assert wb["Budget"]["D2"].value == "=B2+C2"
    assert wb["Budget"]["B2"].number_format == "$#,##0"
    assert wb["Budget"]["A1"].font.b
    assert wb["Budget"].column_dimensions["A"].width == pytest.approx(18.7109375, abs=1)
    assert set(wb.defined_names) == {"Rent", "Food"}
    assert wb["Forecast"]["B7"].value == 1900

    assert charts(repo.book, "Budget") == ["Spending 2026"]
    assert charts(repo.book, "Notes 2026") == ["Q1 only"]
    assert charts(repo.book, "Forecast") == ["Cash forecast"]
    assert "comment A1: 'check with landlord'" in objects(repo.book)["Notes 2026"]


def test_same_cell_and_same_chart_conflict(repo):
    repo.commit("carol", "main", b2=1200, title="Costs")
    result = repo.merge("carol", "alice")
    assert result.returncode != 0
    assert "CONFLICT" in result.stdout + result.stderr
    keep(repo.book, "merged_conflict.xlsx")
    assert_opens_clean(repo.book)

    wb = load_workbook(repo.book)
    assert wb["Budget"]["B2"].value == 1200  # ours kept
    rows = [[c.value for c in r] for r in wb["_merge_conflicts"].iter_rows()]
    assert rows[0][:5] == ["sheet", "cell", "base", "ours (kept)", "theirs"]
    assert ["Budget", "B2", "1000", "1200", "1100"] == rows[1][:5]
    assert rows[1][5].startswith("=HYPERLINK(")
    assert any("chart" in str(r[1]) for r in rows[1:])
    assert charts(repo.book, "Budget") == ["Costs"]
    assert wb.sheetnames[:2] == ["Budget", "Notes 2026"]  # the rename still merged


def test_their_chart_deletion_merges(repo):
    repo.commit("dave", "main", budget_chart=False)
    result = repo.merge("bob", "dave")
    assert result.returncode == 0, result.stdout + result.stderr
    keep(repo.book, "merged_chart_deleted.xlsx")
    assert_opens_clean(repo.book)
    assert charts(repo.book, "Budget") == []
    assert charts(repo.book, "Notes") == ["Q1 only"]
    assert charts(repo.book, "Forecast") == ["Cash forecast"]


def test_renumbered_charts_still_match(tmp_path):
    """We add a chart on an earlier sheet, so the writer saves our existing
    Notes chart as chart2.xml instead of chart1.xml. Their edit to that chart
    must still land on it, with no false conflict."""
    r = Repo(tmp_path)
    r.commit("main", "main", budget_chart=False, notes_chart=True)
    r.commit("ours", "main", budget_chart=True, notes_chart=True)
    r.commit("theirs", "main", budget_chart=False, notes_chart=True, notes_title="Q1 actuals")
    with zipfile.ZipFile(r.book) as z:
        assert "xl/charts/chart2.xml" not in z.namelist()  # theirs: one chart, numbered 1
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stdout + result.stderr
    keep(r.book, "merged_renumbered.xlsx")
    assert_opens_clean(r.book)
    assert charts(r.book, "Budget") == ["Spending"]
    assert charts(r.book, "Notes") == ["Q1 actuals"]


def test_their_comment_edit_merges(repo):
    repo.commit("erin", "main", note_text="landlord says OK")
    result = repo.merge("bob", "erin")
    assert result.returncode == 0, result.stdout + result.stderr
    assert "comment A1: 'landlord says OK'" in objects(repo.book)["Notes"]
    assert charts(repo.book, "Notes") == ["Q1 only"]


def test_excel_namespace_prefixes_survive(tmp_path):
    """Excel marks sheets with mc:Ignorable="x14ac ..."; if a rewrite renamed
    those prefixes Excel would reject the file."""
    paths = {}
    for name, variant in (("base", {}), ("ours", dict(c3=400)), ("theirs", dict(b2=1100))):
        p = tmp_path / f"{name}.xlsx"
        build(p, **variant)
        with zipfile.ZipFile(p) as z:
            parts = {i.filename: z.read(i) for i in z.infolist()}
        parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            b'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac"'
        ).replace(b'<row r="2"', b'<row x14ac:dyDescent="0.25" r="2"')
        with zipfile.ZipFile(p, "w") as z:
            for n, d in parts.items():
                z.writestr(n, d)
        paths[name] = p
    assert xlgit.merge(str(paths["base"]), str(paths["ours"]), str(paths["theirs"])) == 0
    with zipfile.ZipFile(paths["ours"]) as z:
        sheet = z.read("xl/worksheets/sheet1.xml")
    assert b'mc:Ignorable="x14ac"' in sheet and b"x14ac:dyDescent" in sheet
    assert b"ns0:" not in sheet
    assert load_workbook(paths["ours"])["Budget"]["B2"].value == 1100


def test_diff_reports_chart_changes(tmp_path, capsys):
    build(tmp_path / "a.xlsx")
    build(tmp_path / "b.xlsx", **ALICE)
    xlgit.diff(str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx"))
    out = capsys.readouterr().out
    assert "changed        Budget!B2  1000 -> 1100" in out
    # the retitled chart is one change, not a removal plus an addition
    assert "object changed Budget  chart: barChart 'Spending'" in out and "-> chart: barChart 'Spending 2026'" in out
    assert "object removed Budget" not in out and "object added   Budget" not in out
    # the renamed sheet is a rename, with nothing else about it changed
    assert "sheet renamed  Notes -> Notes 2026" in out
    assert "Notes 2026!" not in out and "comment" not in out


def test_diff_says_when_nothing_changed(tmp_path, capsys):
    build(tmp_path / "a.xlsx")
    assert xlgit.diff(str(tmp_path / "a.xlsx"), str(tmp_path / "a.xlsx")) == 0
    assert "No cell or object changes" in capsys.readouterr().out
