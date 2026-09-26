"""Tables and pivot tables across branches.

XlsxWriter writes tables but not pivot tables, so pivots are added to its
output by hand, part by part, the way Excel lays them out: a cache
definition + records per cache, a pivot table definition per pivot, and
the pivot's rendered cells on the sheet.
"""
import zipfile

import pytest
import xlsxwriter
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

import xlgit
from test_merge import Repo, assert_opens_clean, keep, objects

MAIN, REL = xlgit.MAIN, xlgit.REL
PKG_REL, CT = xlgit.PKG_REL, xlgit.CTYPES
ROWS = [("East", 10), ("West", 20), ("East", 30), ("West", 40)]
PIVOT = dict(sheet="Summary", name="PivotTable1", at="A3", func="sum", style="PivotStyleLight16", cache=1)


def aggregate(rows, func):
    regions = sorted({r[0] for r in rows})
    agg = lambda xs: sum(xs) if func == "sum" else sum(xs) / len(xs)
    return regions, [agg([r[1] for r in rows if r[0] == g]) for g in regions], agg([r[1] for r in rows])


def build_sales(path, rows=ROWS, *, cols=("Region", "Amount"), pivots=(PIVOT,), tables=(), formulas=None):
    wb = xlsxwriter.Workbook(str(path))
    sheets = {"Sales": wb.add_worksheet("Sales"), "Summary": wb.add_worksheet("Summary")}
    for spec in list(pivots) + list(tables):
        if spec["sheet"] not in sheets:
            sheets[spec["sheet"]] = wb.add_worksheet(spec["sheet"])
    sheets["Sales"].add_table(0, 0, len(rows), len(cols) - 1, {
        "name": "SalesTbl", "columns": [{"header": c} for c in cols], "data": [list(r) for r in rows]})
    for t in tables:
        sheets[t["sheet"]].add_table(t["ref"], {"columns": [{"header": h} for h in t["header"]], "data": t["data"]})
    for p in pivots:
        ws = sheets[p["sheet"]]
        col, row = coordinate_from_string(p["at"])
        c, r = column_index_from_string(col) - 1, row - 1
        regions, values, total = aggregate(rows, p["func"])
        ws.write(r, c, "Row Labels")
        ws.write(r, c + 1, f"{'Sum' if p['func'] == 'sum' else 'Average'} of Amount")
        for i, (g, v) in enumerate(zip(regions, values), start=1):
            ws.write(r + i, c, g)
            ws.write_number(r + i, c + 1, v)
        ws.write(r + len(regions) + 1, c, "Grand Total")
        ws.write_number(r + len(regions) + 1, c + 1, total)
    for sheet, cells in (formulas or {}).items():
        for cell, f in cells.items():
            if f.startswith("{"):  # an array formula
                sheets[sheet].write_array_formula(f"{cell}:{cell}", f)
            else:
                sheets[sheet].write_formula(cell, f)
    wb.close()
    if pivots:
        add_pivots(path, rows, cols, pivots)


def add_pivots(path, rows, cols, pivots):
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i) for i in z.infolist()}
    x = lambda name: etree.fromstring(parts[name])
    wb, wb_rels, ct = x("xl/workbook.xml"), x("xl/_rels/workbook.xml.rels"), x("[Content_Types].xml")
    sheet_part = {s.get(xlgit.RID): s.get("name") for s in wb.find(f"{{{MAIN}}}sheets")}
    sheet_part = {sheet_part[r.get("Id")]: "xl/" + r.get("Target") for r in wb_rels if r.get("Id") in sheet_part}

    def rel(root, typ, target):
        rid = f"rId{len(root) + 1}"
        etree.SubElement(root, f"{{{PKG_REL}}}Relationship", Id=rid, Type=f"{REL}/{typ}", Target=target)
        return rid

    def override(name, kind):
        etree.SubElement(ct, f"{{{CT}}}Override", PartName="/" + name,
                         ContentType=f"application/vnd.openxmlformats-officedocument.spreadsheetml.{kind}+xml")

    regions = sorted({r[0] for r in rows})
    amounts = [r[1] for r in rows]
    pcs = etree.SubElement(wb, f"{{{MAIN}}}pivotCaches")
    xlgit.Merger.place(wb, pcs, xlgit.WB_ORDER)
    for cid in sorted({p["cache"] for p in pivots}):
        extra_fields = "".join(f'<cacheField name="{c}" numFmtId="0"><sharedItems/></cacheField>' for c in cols[2:])
        parts[f"xl/pivotCache/pivotCacheDefinition{cid}.xml"] = f"""<pivotCacheDefinition xmlns="{MAIN}" xmlns:r="{REL}"
 r:id="rId1" refreshedBy="test" refreshedDate="{45000 + sum(amounts)}" createdVersion="6" refreshedVersion="6"
 minRefreshableVersion="3" recordCount="{len(rows)}">
<cacheSource type="worksheet"><worksheetSource ref="A1:{get_column_letter(len(cols))}{len(rows) + 1}" sheet="Sales"/></cacheSource>
<cacheFields count="{len(cols)}">
<cacheField name="{cols[0]}" numFmtId="0"><sharedItems count="{len(regions)}">{"".join(f'<s v="{g}"/>' for g in regions)}</sharedItems></cacheField>
<cacheField name="{cols[1]}" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1" containsInteger="1" minValue="{min(amounts)}" maxValue="{max(amounts)}"/></cacheField>
{extra_fields}</cacheFields></pivotCacheDefinition>""".encode()
        records = "".join(f'<r><x v="{regions.index(r[0])}"/><n v="{r[1]}"/>'
                          + "".join(f'<s v="{v}"/>' for v in r[2:]) + "</r>" for r in rows)
        parts[f"xl/pivotCache/pivotCacheRecords{cid}.xml"] = (
            f'<pivotCacheRecords xmlns="{MAIN}" xmlns:r="{REL}" count="{len(rows)}">{records}</pivotCacheRecords>').encode()
        crels = etree.Element(f"{{{PKG_REL}}}Relationships", nsmap={None: PKG_REL})
        rel(crels, "pivotCacheRecords", f"pivotCacheRecords{cid}.xml")
        parts[f"xl/pivotCache/_rels/pivotCacheDefinition{cid}.xml.rels"] = etree.tostring(crels)
        etree.SubElement(pcs, f"{{{MAIN}}}pivotCache", cacheId=str(cid)).set(
            xlgit.RID, rel(wb_rels, "pivotCacheDefinition", f"pivotCache/pivotCacheDefinition{cid}.xml"))
        override(f"xl/pivotCache/pivotCacheDefinition{cid}.xml", "pivotCacheDefinition")
        override(f"xl/pivotCache/pivotCacheRecords{cid}.xml", "pivotCacheRecords")

    for i, p in enumerate(pivots, start=1):
        col, row = coordinate_from_string(p["at"])
        c = column_index_from_string(col)
        ref = f"{col}{row}:{get_column_letter(c + 1)}{row + len(regions) + 1}"
        subtotal = "" if p["func"] == "sum" else ' subtotal="average"'
        caption = f"{'Sum' if p['func'] == 'sum' else 'Average'} of Amount"
        items = "".join(f'<item x="{k}"/>' for k in range(len(regions)))
        row_items = "<i><x/></i>" + "".join(f'<i><x v="{k}"/></i>' for k in range(1, len(regions)))
        extra = '<pivotField showAll="0"/>' * (len(cols) - 2)
        name = f"xl/pivotTables/pivotTable{i}.xml"
        parts[name] = f"""<pivotTableDefinition xmlns="{MAIN}" name="{p['name']}" cacheId="{p['cache']}"
 applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0"
 applyAlignmentFormats="0" applyWidthHeightFormats="1" dataCaption="Values" updatedVersion="6"
 minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1" createdVersion="6" indent="0"
 outline="1" outlineData="1" multipleFieldFilters="0">
<location ref="{ref}" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/>
<pivotFields count="{len(cols)}"><pivotField axis="axisRow" showAll="0"><items count="{len(regions) + 1}">{items}<item t="default"/></items></pivotField>
<pivotField dataField="1" showAll="0"/>{extra}</pivotFields>
<rowFields count="1"><field x="0"/></rowFields>
<rowItems count="{len(regions) + 1}">{row_items}<i t="grand"><x/></i></rowItems>
<colItems count="1"><i/></colItems>
<dataFields count="1"><dataField name="{caption}" fld="1"{subtotal} baseField="0" baseItem="0"/></dataFields>
<pivotTableStyleInfo name="{p['style']}" showRowHeaders="1" showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/>
</pivotTableDefinition>""".encode()
        prels = etree.Element(f"{{{PKG_REL}}}Relationships", nsmap={None: PKG_REL})
        rel(prels, "pivotCacheDefinition", f"../pivotCache/pivotCacheDefinition{p['cache']}.xml")
        parts[f"xl/pivotTables/_rels/pivotTable{i}.xml.rels"] = etree.tostring(prels)
        override(name, "pivotTable")
        spart = sheet_part[p["sheet"]]
        srels_name = xlgit.rels_name(spart)
        srels = x(srels_name) if srels_name in parts else etree.Element(f"{{{PKG_REL}}}Relationships", nsmap={None: PKG_REL})
        rel(srels, "pivotTable", f"../pivotTables/pivotTable{i}.xml")
        parts[srels_name] = etree.tostring(srels)

    parts["xl/workbook.xml"], parts["xl/_rels/workbook.xml.rels"] = etree.tostring(wb), etree.tostring(wb_rels)
    parts["[Content_Types].xml"] = etree.tostring(ct)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def repo_with(tmp_path, base, ours, theirs):
    r = Repo(tmp_path)
    for branch, spec in (("main", base), ("ours", ours), ("theirs", theirs)):
        r.git("checkout", "-q", "main", check=branch != "main")
        if branch != "main":
            r.git("checkout", "-q", "-b", branch)
        build_sales(r.book, **spec)
        r.git("add", "-A")
        r.git("commit", "-q", "-m", branch)
    return r


def tables(path):
    return [o for items in objects(path).values() for o in items if o.startswith("table:")]


def pivots(path):
    """{sheet: [(name, location, values)]} straight from the pivot XML."""
    pkg = xlgit.Package.open(str(path))
    out = {}
    for sheet, _, part in pkg.sheets():
        for _, typ, t, ext in pkg.rels(part):
            if xlgit.kind_of(typ) == "pivotTable":
                p = pkg.xml(t)
                out.setdefault(sheet, []).append(
                    (p.get("name"), p.find(f"{{{MAIN}}}location").get("ref"),
                     [d.get("name") for d in p.iter(f"{{{MAIN}}}dataField")]))
    return out


def pivot_integrity(path):
    """Each pivot's cacheId must be registered in workbook.xml against the
    cache it links to, and every registered cache must exist."""
    pkg = xlgit.Package.open(str(path))
    wb = pkg.xml(pkg.workbook)
    targets = {rid: t for rid, _, t, _ in pkg.rels(pkg.workbook)}
    registered = {pc.get("cacheId"): targets[pc.get(xlgit.RID)] for pc in wb.iter(f"{{{MAIN}}}pivotCache")}
    assert all(c in pkg.parts for c in registered.values())
    caches = set()
    for sheet, _, part in pkg.sheets():
        for _, typ, t, ext in pkg.rels(part):
            if xlgit.kind_of(typ) == "pivotTable":
                cache = pkg.target(t, "pivotCacheDefinition")
                assert registered.get(pkg.xml(t).get("cacheId")) == cache
                caches.add(cache)
    return pkg, caches


def test_table_rows_and_column_merge(tmp_path):
    r = repo_with(tmp_path, dict(pivots=()),
                  dict(rows=ROWS + [("North", 50)], pivots=()),
                  dict(cols=("Region", "Amount", "Rep"), rows=[x + (f"rep{i}",) for i, x in enumerate(ROWS)], pivots=()))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    keep(r.book, "tables_rows_and_column.xlsx")
    assert_opens_clean(r.book)
    assert tables(r.book) == ["table: SalesTbl A1:C6 columns=['Region', 'Amount', 'Rep']"]
    ws = load_workbook(r.book)["Sales"]
    assert [ws[f"A6"].value, ws["B6"].value, ws["C1"].value, ws["C2"].value] == ["North", 50, "Rep", "rep0"]


def test_both_sides_add_a_table_named_table2(tmp_path):
    ours_t = dict(sheet="Sales", ref="E1:F3", header=["K", "V"], data=[["a", 1], ["b", 2]])
    theirs_t = dict(sheet="Summary", ref="H1:I3", header=["X", "Y"], data=[["c", 3], ["d", 4]])
    r = repo_with(tmp_path, dict(pivots=()), dict(pivots=(), tables=[ours_t]),
                  dict(pivots=(), tables=[theirs_t], formulas={"Summary": {"K1": "=SUM(Table2[Y])"}}))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    assert "renamed to Table3" in result.stderr
    keep(r.book, "tables_both_added.xlsx")
    assert_opens_clean(r.book)
    assert sorted(tables(r.book)) == ["table: SalesTbl A1:B5 columns=['Region', 'Amount']",
                                      "table: Table2 E1:F3 columns=['K', 'V']",
                                      "table: Table3 H1:I3 columns=['X', 'Y']"]
    pkg = xlgit.Package.open(str(r.book))
    ids = [pkg.xml(t).get("id") for t in pkg.parts if t.startswith("xl/tables/")]
    assert len(ids) == len(set(ids)) == 3
    assert load_workbook(r.book)["Summary"]["K1"].value == "=SUM(Table3[Y])"


def test_their_table_deletion_merges(tmp_path):
    extra = dict(sheet="Summary", ref="H1:I3", header=["X", "Y"], data=[["c", 3], ["d", 4]])
    r = repo_with(tmp_path, dict(pivots=(), tables=[extra]),
                  dict(pivots=(), tables=[extra], rows=[("East", 11)] + ROWS[1:]),
                  dict(pivots=()))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    assert_opens_clean(r.book)
    assert tables(r.book) == ["table: SalesTbl A1:B5 columns=['Region', 'Amount']"]
    assert load_workbook(r.book)["Sales"]["B2"].value == 11


def test_pivot_our_refresh_their_layout_change(tmp_path):
    """We change source data (our pivot refreshed), they switch the pivot to
    Average. Result: their pivot layout, our data, refresh on open."""
    r = repo_with(tmp_path, {}, dict(rows=[("East", 15)] + ROWS[1:]), dict(pivots=(dict(PIVOT, func="avg"),)))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    keep(r.book, "pivot_refresh_vs_layout.xlsx")
    assert_opens_clean(r.book)
    pkg, caches = pivot_integrity(r.book)
    assert pivots(r.book) == {"Summary": [("PivotTable1", "A3:B6", ["Average of Amount"])]}
    assert load_workbook(r.book)["Sales"]["B2"].value == 15
    assert all(pkg.xml(c).get("refreshOnLoad") == "1" for c in caches)


def test_pivot_layout_changed_on_both_sides_conflicts(tmp_path):
    r = repo_with(tmp_path, {}, dict(pivots=(dict(PIVOT, style="PivotStyleMedium9"),)),
                  dict(pivots=(dict(PIVOT, func="avg"),)))
    result = r.merge("ours", "theirs")
    assert result.returncode != 0
    assert "(pivot)" in result.stderr
    assert_opens_clean(r.book)
    pivot_integrity(r.book)
    assert pivots(r.book) == {"Summary": [("PivotTable1", "A3:B6", ["Sum of Amount"])]}
    assert b"PivotStyleMedium9" in xlgit.Package.open(str(r.book)).parts["xl/pivotTables/pivotTable1.xml"]


def test_their_new_pivot_shares_the_cache(tmp_path):
    second = dict(PIVOT, name="PivotTable2", at="E3")
    r = repo_with(tmp_path, {}, dict(rows=[("East", 15)] + ROWS[1:]), dict(pivots=(PIVOT, second)))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    keep(r.book, "pivot_new_shared_cache.xlsx")
    assert_opens_clean(r.book)
    pkg, caches = pivot_integrity(r.book)
    assert len(caches) == 1
    assert [p[0] for p in pivots(r.book)["Summary"]] == ["PivotTable1", "PivotTable2"]
    assert load_workbook(r.book)["Summary"]["E3"].value == "Row Labels"


def test_their_new_sheet_with_its_own_pivot(tmp_path):
    other = dict(PIVOT, sheet="Summary2", cache=2)
    r = repo_with(tmp_path, {}, dict(rows=[("East", 15)] + ROWS[1:]), dict(pivots=(PIVOT, other)))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    keep(r.book, "pivot_new_sheet.xlsx")
    assert_opens_clean(r.book)
    pkg, caches = pivot_integrity(r.book)
    assert len(caches) == 2
    assert set(pivots(r.book)) == {"Summary", "Summary2"}


def test_their_pivot_deletion_merges(tmp_path):
    second = dict(PIVOT, name="PivotTable2", at="E3")
    r = repo_with(tmp_path, dict(pivots=(PIVOT, second)),
                  dict(pivots=(PIVOT, second), rows=[("East", 15)] + ROWS[1:]), dict(pivots=(PIVOT,)))
    result = r.merge("ours", "theirs")
    assert result.returncode == 0, result.stderr
    keep(r.book, "pivot_deleted.xlsx")
    assert_opens_clean(r.book)
    pivot_integrity(r.book)
    assert [p[0] for p in pivots(r.book)["Summary"]] == ["PivotTable1"]


def test_diff_lists_tables_and_pivots(tmp_path, capsys):
    build_sales(tmp_path / "a.xlsx")
    build_sales(tmp_path / "b.xlsx", pivots=(dict(PIVOT, func="avg"),))
    xlgit.diff(str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx"))
    out = capsys.readouterr().out
    assert "pivot table: PivotTable1 at A3:B6 rows=['Region'] cols=[] values=['Sum of Amount']" in out
    assert "values=['Average of Amount']" in out
