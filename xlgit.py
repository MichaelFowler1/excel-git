# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: Apache-2.0
"""xlgit: make Excel workbooks behave like code in git and GitHub.

Run `xlgit` for help. Set up once with `xlgit install`; after that git diff
and git merge understand .xlsx/.xlsm files in every repository.

The merge edits the workbook's XML parts in place instead of re-saving it
through a spreadsheet library, so charts, images, formatting, comments,
tables, pivot tables and macros in OUR copy survive untouched, and their
edits to those objects are carried over.
"""
import copy
import datetime
import hashlib
import io
import os
import posixpath
import re
import subprocess
import sys
import zipfile
from collections import namedtuple

from lxml import etree
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter
from openpyxl.utils.datetime import (CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel, from_ISO8601,
                                     to_excel)

__version__ = "0.1.0"

EXTS = ("*.xlsx", "*.xlsm")

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CTYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DRAW = "http://schemas.openxmlformats.org/drawingml/2006/main"
RID = f"{{{REL}}}id"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
WORKSHEET_REL = REL + "/worksheet"
REL_TABLE = REL + "/table"
REL_PIVOT = REL + "/pivotTable"
REL_CACHE = REL + "/pivotCacheDefinition"
WORKSHEET_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
CONFLICT_SHEET = "_merge_conflicts"
# Sheet parts holding cells: worksheets and Excel 4 macro sheets.
SHEET_ROOTS = ("worksheet", "macrosheet")

# Worksheet child order from the OOXML schema; new elements must slot in here.
WS_ORDER = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData",
            "sheetCalcPr", "sheetProtection", "protectedRanges", "scenarios", "autoFilter",
            "sortState", "dataConsolidate", "customSheetViews", "mergeCells", "phoneticPr",
            "conditionalFormatting", "dataValidations", "hyperlinks", "printOptions",
            "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
            "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing",
            "legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects",
            "controls", "webPublishItems", "tableParts", "extLst"]
WB_ORDER = ["fileVersion", "fileSharing", "workbookPr", "workbookProtection", "bookViews",
            "sheets", "functionGroups", "externalReferences", "definedNames", "calcPr",
            "oleSize", "customWorkbookViews", "pivotCaches", "smartTagPr", "smartTagTypes",
            "webPublishing", "fileRecoveryPr", "webPublishObjects", "extLst"]
TABLE_ORDER = ["autoFilter", "sortState", "tableColumns", "tableStyleInfo", "extLst"]

# Parts that are merged some other way (cells, sheet list, names) or that
# every save rewrites. Anything else their branch adds that doesn't get
# merged is reported instead of silently dropped.
NOT_OBJECT_PART = re.compile(
    r"^(\[Content_Types\]\.xml|docProps/.*|customXml/.*|xl/workbook\.xml|xl/worksheets/.*"
    r"|xl/sharedStrings\.xml|xl/calcChain\.xml|xl/styles\.xml|xl/theme/.*|xl/metadata\.xml"
    r"|xl/persons/.*|xl/printerSettings/.*|.*\.rels)$")
# Workbook-level parts left as ours.
WB_KEEP = {"sharedStrings", "styles", "theme", "calcChain", "metadata", "sheetMetadata", "person"}
# Pivot cache attributes that change on every refresh without meaning anything.
VOLATILE_CACHE_ATTRS = ("refreshedDate", "refreshedDateIso", "refreshedBy", "refreshedVersion",
                        "refreshOnLoad", "recordCount", "invalid")

ArrayF = namedtuple("ArrayF", "ref text")
Opaque = namedtuple("Opaque", "desc")


# Workbooks come from anyone who can open a pull request: never resolve
# entities (external ones read local files; older lxml resolved them by
# default) or touch the network.
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


def parse_xml(data):
    return etree.fromstring(data, _PARSER)


def m(tag):
    return f"{{{MAIN}}}{tag}"


def local(el):
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def kind_of(rel_type):
    return rel_type.rsplit("/", 1)[-1]


def digest(data):
    return hashlib.sha1(data or b"").hexdigest()[:8]


def parse_ref(ref):
    a, _, b = ref.replace("$", "").partition(":")
    (c1, r1), (c2, r2) = coordinate_from_string(a), coordinate_from_string(b or a)
    return column_index_from_string(c1), r1, column_index_from_string(c2), r2


def format_ref(c1, r1, c2, r2):
    lo, hi = f"{get_column_letter(c1)}{r1}", f"{get_column_letter(c2)}{r2}"
    return lo if lo == hi else f"{lo}:{hi}"


def overlaps(a, b):
    a1, ar1, a2, ar2 = parse_ref(a)
    b1, br1, b2, br2 = parse_ref(b)
    return not (a2 < b1 or b2 < a1 or ar2 < br1 or br2 < ar1)


_QUOTED = re.compile(r"""('(?:[^']|'')*'|"(?:[^"]|"")*")""")


def rename_refs(text, renames):
    """Point structured references (Table2[Amount], =SUM(Table2)) at renamed
    tables. Quoted sheet names ('Table2 notes'!A1) and text ("Table2") are
    left alone: they only look like the table's name."""
    parts = _QUOTED.split(text)
    for i in range(0, len(parts), 2):  # even pieces are outside quotes
        for old, new in renames.items():
            parts[i] = re.sub(rf"(?<![\w.]){re.escape(old)}(?=\[|(?![\w.(!]))", new, parts[i], flags=re.I)
    return "".join(parts)


def xdecode(text):
    """Undo OOXML's _xHHHH_ escaping (a line break in a table column name is
    stored as _x000a_)."""
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda mt: chr(int(mt.group(1), 16)), text or "")


def xencode(text):
    text = re.sub(r"_(?=x[0-9A-Fa-f]{4}_)", "_x005F_", text)
    return re.sub(r"[\x00-\x1f]", lambda mt: f"_x{ord(mt.group(0)):04x}_", text)


def strip_dxf(root):
    """Differential formats (dxfId) and custom table/pivot styles index into
    styles.xml; when the branches' styles differ those numbers are meaningless."""
    for el in root.iter():
        for a in list(el.attrib):
            if a.lower().endswith("dxfid"):
                del el.attrib[a]
    for el in root.iter(m("tableStyleInfo"), m("pivotTableStyleInfo")):
        name = el.get("name") or ""
        if not re.match(r"(Table|Pivot)Style(Light|Medium|Dark)\d+$", name):
            el.set("name", "TableStyleMedium2" if local(el) == "tableStyleInfo" else "PivotStyleLight16")


# ---------- package (zip of XML parts) helpers ----------

def rels_name(owner):
    if owner is None:  # a sheet with no part behind it (Excel 4 macro sheets)
        return None
    if owner == "":
        return "_rels/.rels"
    d, b = posixpath.split(owner)
    return f"{d}/_rels/{b}.rels" if d else f"_rels/{b}.rels"


def resolve(owner, target):
    if target.startswith("/"):
        return target[1:]
    return posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))


def relative(target, owner):
    d = posixpath.dirname(owner)
    return posixpath.relpath(target, d) if d else target


def reachable(parts):
    """Every part reachable from the package root through relationships."""
    seen, stack = set(), [""]
    while stack:
        owner = stack.pop()
        rn = rels_name(owner)
        if rn not in parts:
            continue
        for r in parse_xml(parts[rn]):
            if r.get("TargetMode") == "External":
                continue
            t = resolve(owner, r.get("Target"))
            if t in parts and t not in seen:
                seen.add(t)
                stack.append(t)
    return seen


class Package:
    """Read-only raw view of an .xlsx zip."""

    def __init__(self, data=b""):
        self.parts, self.order = {}, []
        self._ro = {}
        if data:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for info in z.infolist():
                    if not info.is_dir():
                        self.parts[info.filename] = z.read(info)
                        self.order.append(info.filename)
        self.workbook = next((t for _, typ, t, _ in self.rels("") if typ.endswith("/officeDocument")),
                             "xl/workbook.xml")

    @classmethod
    def open(cls, path):
        if not path or not os.path.exists(path) or os.path.getsize(path) == 0:
            return cls()
        with open(path, "rb") as f:
            return cls(f.read())

    def xml(self, name):
        return parse_xml(self.parts[name]) if name in self.parts else None

    def scan_sheet(self, part, sst=None, dates=frozenset(), styles_for=None):
        """Stream through a worksheet part without keeping it in memory.
        Returns (cells, children, styles): cell values (if sst is given), the
        worksheet's top-level elements other than sheetData, and the style of
        each cell in styles_for. Sheets can hold millions of empty formatted
        cells; a parsed copy of one can take gigabytes."""
        cells, children, styles, shared = {}, [], {}, {}
        root_tag = None
        ROW, C, SD = m("row"), m("c"), m("sheetData")
        rnum = 0
        depth = 0
        for event, el in etree.iterparse(io.BytesIO(self.parts[part]), events=("start", "end"),
                                         resolve_entities=False, no_network=True):
            if event == "start":
                depth += 1
                if depth == 1:
                    root_tag = local(el)
                    if root_tag not in SHEET_ROOTS:
                        break
                continue
            depth -= 1
            if el.tag == ROW and depth == 2:
                r = el.get("r")
                rnum = int(r) if r else rnum + 1
                prev = None
                for c in el.iterchildren(C):
                    coord = c.get("r")
                    if coord is None:
                        if isinstance(prev, str):
                            prev = column_index_from_string(coordinate_from_string(prev)[0])
                        prev = (prev or 0) + 1
                        coord = f"{get_column_letter(prev)}{rnum}"
                    else:
                        prev = coord
                    if sst is not None:
                        v = _cell_value(c, coord, sst, shared, dates)
                        if v is not None:
                            cells[coord] = v
                    if styles_for is not None and coord in styles_for:
                        styles[coord] = c.get("s")
                el.clear()
                while el.getprevious() is not None:
                    del el.getparent()[0]
            elif depth == 1:
                if el.tag == SD:
                    el.clear()
                else:
                    children.append(el)
        return cells, children, styles

    def sheet_children(self, part):
        """The worksheet's top-level elements except its cells (drawing,
        legacyDrawing, tableParts...), cached."""
        key = ("children", part)
        if key not in self._ro:
            self._ro[key] = self.scan_sheet(part)[1] if part in self.parts else []
        return self._ro[key]

    def ro(self, name):
        """Parsed part, cached: for reading only, never modify it."""
        if name not in self._ro:
            self._ro[name] = self.xml(name)
        return self._ro[name]

    def rels(self, owner):
        """[(rId, type, target, external)]; internal targets are absolute part names."""
        root = self.xml(rels_name(owner))
        if root is None:
            return []
        out = []
        for r in root:
            ext = r.get("TargetMode") == "External"
            out.append((r.get("Id"), r.get("Type"), r.get("Target") if ext else resolve(owner, r.get("Target")), ext))
        return out

    def target(self, owner, kind):
        return next((t for _, typ, t, ext in self.rels(owner) if not ext and kind_of(typ) == kind), None)

    def sheets(self):
        """[(name, sheetId, part)] in workbook order."""
        wb = self.xml(self.workbook)
        if wb is None:
            return []
        targets = {rid: t for rid, _, t, _ in self.rels(self.workbook)}
        return [(s.get("name"), s.get("sheetId"), targets.get(s.get(RID)))
                for s in wb.find(m("sheets"))]

    def content_type(self, name):
        ct = self.xml("[Content_Types].xml")
        for o in ct.iter(f"{{{CTYPES}}}Override"):
            if o.get("PartName").lstrip("/").lower() == name.lower():
                return ("override", o.get("ContentType"))
        ext = name.rsplit(".", 1)[-1].lower()
        for d in ct.iter(f"{{{CTYPES}}}Default"):
            if d.get("Extension").lower() == ext:
                return ("default", ext, d.get("ContentType"))
        return None

    def shared_strings(self):
        t = self.target(self.workbook, "sharedStrings")
        return list(self.xml(t).iter(m("si"))) if t in self.parts else []

    def styles_bytes(self):
        return self.parts.get(self.target(self.workbook, "styles"))

    def sheet_object(self, part, key):
        """Target part of a sheet's drawing / legacyDrawing (element-wired) or comments (rel-only)."""
        if not part or part not in self.parts:
            return None, None
        rels = {rid: (typ, t) for rid, typ, t, ext in self.rels(part) if not ext}
        if key == "comments":
            return next(((t, typ) for typ, t in rels.values() if kind_of(typ) == "comments"), (None, None))
        el = next((c for c in self.sheet_children(part) if local(c) == key), None)
        if el is None:
            return None, None
        typ, t = rels.get(el.get(RID), (None, None))
        return t, typ


class Keys:
    """Structural identity for every part of a package.

    Excel and the Python writers renumber files on every save (add a chart
    to an early sheet and every later chart1.xml becomes chart2.xml), so
    part names can't be used to find "the same chart" in another version.
    Instead each part gets a key describing what it is: the drawing of sheet
    Budget, the shape called "Chart 1" in that drawing, table id 3, the
    pivot table called PivotTable1 on sheet Summary, and so on."""

    def __init__(self, pkg, sheet_keys, base_keys=()):
        self.pkg = pkg
        self.sheet_keys = sheet_keys
        self.key_of, self.part_of = {}, {}
        self._xml = {}
        self.pivots = []
        for _, typ, t, ext in pkg.rels(""):
            if not ext and t in pkg.parts:
                self._set(t, ("workbook",) if kind_of(typ) == "officeDocument" else ("pkg", kind_of(typ)))
        wb = pkg.workbook
        if wb in pkg.parts:
            sheet_rids = {s.get(RID): s.get("name") for s in pkg.xml(wb).find(m("sheets"))}
            for rid, typ, t, ext in pkg.rels(wb):
                if ext or t not in pkg.parts or kind_of(typ) == "pivotCacheDefinition":
                    continue
                if rid in sheet_rids:
                    self._set(t, ("sheet", sheet_keys.get(sheet_rids[rid], sheet_rids[rid])))
                else:
                    self._set(t, ("wb", kind_of(typ)))
            for part in list(self.key_of):
                self._walk(part)
            # A pivot cache is known by the data it summarizes; if several
            # caches read the same source, by the pivot tables that use them.
            users = {}
            for pkey, ppart in self.pivots:
                cache = pkg.target(ppart, "pivotCacheDefinition")
                if cache in pkg.parts:
                    users.setdefault(cache, []).append(pkey)
            # Caches whose pivots existed in base claim their keys first, so a
            # new cache on the same data doesn't steal an old one's identity.
            rank = lambda cache: (not any(k in base_keys for k in users[cache]), min(repr(k) for k in users[cache]))
            for cache in sorted(users, key=rank):
                self._set(cache, ("cache", self._cache_source(cache)))
                self._walk(cache)
            for _, typ, t, ext in pkg.rels(wb):
                if not ext and kind_of(typ) == "pivotCacheDefinition" and t in pkg.parts and t not in self.key_of:
                    self._set(t, ("cache", "unused", t))
                    self._walk(t)
        for name in pkg.parts:
            if name not in self.key_of and not name.endswith(".rels"):
                self._set(name, ("part", name))

    def _set(self, part, key):
        if part in self.key_of:
            return
        n = 2
        base = key
        while key in self.part_of:
            key = base + (n,)
            n += 1
        self.key_of[part] = key
        self.part_of[key] = part

    def _tree(self, part):
        if part not in self._xml:
            try:
                self._xml[part] = parse_xml(self.pkg.parts[part])
            except etree.XMLSyntaxError:
                self._xml[part] = None
        return self._xml[part]

    def _cache_source(self, cache):
        root = self._tree(cache)
        src = root.find(m("cacheSource")) if root is not None else None
        if src is None:
            return ("unknown",)
        ws = src.find(m("worksheetSource"))
        if ws is None:
            return (src.get("type"), src.get("connectionId"))
        if ws.get("name"):
            return ("name", ws.get("name"))
        return ("sheet", self.sheet_keys.get(ws.get("sheet"), ws.get("sheet")))

    def _shape_name(self, drawing, rid):
        """Name of the drawing shape (chart, picture) that uses relationship rid."""
        root = self._tree(drawing)
        for el in (root.iter() if root is not None else []):
            if not any(k.startswith(f"{{{REL}}}") and v == rid for k, v in el.attrib.items()):
                continue
            anc = el
            while anc is not None:
                for child in anc:
                    if local(child).startswith("nv"):
                        for c in child:
                            if local(c) == "cNvPr":
                                return c.get("name") or c.get("id")
                anc = anc.getparent()
        return rid

    def _walk(self, owner):
        okey = self.key_of[owner]
        for rid, typ, t, ext in self.pkg.rels(owner):
            if ext or t not in self.pkg.parts or t in self.key_of:
                continue
            kind = kind_of(typ)
            if kind == "pivotCacheDefinition":
                continue  # keyed through the pivot tables that use it
            if okey[0] == "sheet":
                sk = okey[1]
                if kind == "drawing":
                    key = ("drawing", sk)
                elif kind == "vmlDrawing":
                    el = next((e for top in self.pkg.sheet_children(owner) for e in top.iter()
                               if e.get(RID) == rid), None)
                    key = ("vml", sk, local(el) if el is not None else rid)
                elif kind == "comments":
                    key = ("comments", sk)
                elif kind == "table":
                    key = ("table", (self._tree(t).get("id") if self._tree(t) is not None else t))
                elif kind == "pivotTable":
                    key = ("pivot", sk, self._tree(t).get("name") if self._tree(t) is not None else t)
                    self.pivots.append((key, t))
                else:
                    key = ("sub", okey, kind)
            elif okey[0] == "drawing":
                key = ("shape", okey, self._shape_name(owner, rid), kind)
            else:
                key = ("sub", okey, kind)
            self._set(t, key)
            if key[0] == "pivot":
                self.pivots[-1] = (self.key_of[t], t)
            self._walk(t)


# ---------- reading cells ----------

class Text(str):
    """Text that would otherwise read as a formula or an error ("=== Total ===",
    "#N/A" typed as text). Never equal to the formula or error it looks like."""

    def __eq__(self, other):
        return isinstance(other, Text) and str.__eq__(self, other)

    def __ne__(self, other):
        return not self == other

    __hash__ = str.__hash__


class DateNum(float):
    """A number in a date-formatted cell. It merges as the exact number
    stored; the date is only for display."""

    def display(self):
        try:
            return from_excel(float(self)).isoformat(sep=" ").replace(" 00:00:00", "")
        except (OverflowError, ValueError, TypeError):
            return repr(float(self))


def same(a, b):
    """Cell values are equal, with TRUE never equal to 1. Numbers are
    compared to 15 significant digits, the precision Excel works to:
    LibreOffice saves 0.1+0.2 as 0.3 where Excel saves 0.30000000000000004,
    and that isn't a change anyone made."""
    if (type(a) is bool) != (type(b) is bool):
        return False
    if a == b:
        return True
    return isinstance(a, (int, float)) and isinstance(b, (int, float)) and f"{a:.15g}" == f"{b:.15g}"


def same_cells(a, b):
    a, b = a or {}, b or {}
    return a.keys() == b.keys() and all(same(v, b[k]) for k, v in a.items())


def _number(text):
    if "_" not in text:
        try:
            return int(text)
        except ValueError:
            try:
                return float(text)
            except ValueError:
                pass
    return Opaque(f"number {text!r}")


_F, _V, _IS = m("f"), m("v"), m("is")


def _text(el):
    """Plain text of a string item or inline string, skipping phonetic runs."""
    return "".join(t.text or "" for t in el.iter(m("t"))
                   if t.getparent() is None or local(t.getparent()) != "rPh")


def _date_styles(pkg):
    """Indexes of the cell styles (s="...") that show numbers as dates."""
    part = pkg.target(pkg.workbook, "styles")
    root = pkg.xml(part) if part in pkg.parts else None
    xfs = root.find(m("cellXfs")) if root is not None else None
    if xfs is None:
        return set()
    custom = {nf.get("numFmtId"): nf.get("formatCode") or "" for nf in root.iter(m("numFmt"))}
    out = set()
    for i, xf in enumerate(xfs.findall(m("xf"))):
        fid = xf.get("numFmtId") or "0"
        code = custom.get(fid, BUILTIN_FORMATS.get(int(fid)) if fid.isdigit() else None)
        if code and is_date_format(code):
            out.add(str(i))
    return out


def iter_sheet_cells(root):
    """(coord, cell element) for every cell of a worksheet. References are
    optional in the file (a row or cell without one follows the previous)."""
    sd = root.find(m("sheetData")) if root is not None else None
    if sd is None:
        return
    ROW, C = m("row"), m("c")
    rnum = 0
    for row in sd.iterchildren(ROW):
        r = row.get("r")
        rnum = int(r) if r else rnum + 1
        prev = None  # last explicit reference, or the last column number
        for c in row.iterchildren(C):
            coord = c.get("r")
            if coord is None:
                if isinstance(prev, str):
                    prev = column_index_from_string(coordinate_from_string(prev)[0])
                prev = (prev or 0) + 1
                coord = f"{get_column_letter(prev)}{rnum}"
            else:
                prev = coord
            yield coord, c


def sheet_cells(pkg, part, sst, dates=frozenset()):
    """{coord: value} for one worksheet part, read straight from its XML.
    Formulas are '=...' strings (shared formulas spelled out per cell), text
    that looks like a formula or error is Text, errors are '#...' strings."""
    if part not in pkg.parts:
        return {}
    cells, children, _ = pkg.scan_sheet(part, sst, dates)
    pkg._ro[("children", part)] = children
    return cells


def _cell_value(c, coord, sst, shared, dates):
    t = c.get("t", "n")
    f = v = is_ = None
    for ch in c:
        tag = ch.tag
        if tag == _V:
            v = ch
        elif tag == _F:
            f = ch
        elif tag == _IS:
            is_ = ch
    if f is not None:
        ft = f.get("t")
        text = f.text or ""
        if ft == "array":
            return ArrayF(f.get("ref") or coord, "=" + text)
        if ft == "dataTable":
            return Opaque("data table")
        if ft == "shared" and f.get("si") is not None:
            si = f.get("si")
            if text.strip():
                shared[si] = (coord, text, None)
            elif si in shared:
                origin, master, tr = shared[si]
                try:
                    if tr is None:  # parsing the formula is the slow part: once per group
                        tr = Translator("=" + master, origin)
                        shared[si] = (origin, master, tr)
                    return tr.translate_formula(coord)
                except Exception:
                    return Opaque(f"shared formula {master!r}")
        if text:
            return "=" + text
    if t == "inlineStr":
        v = _text(is_) if is_ is not None else None
        return Text(v) if v is not None and (v.startswith("=") or v in ERRORS) else v
    if v is None or v.text is None:
        return None
    if t == "n":
        n = _number(v.text)
        return DateNum(n) if dates and c.get("s") in dates and not isinstance(n, Opaque) else n
    if t == "s":
        try:
            s = sst[int(v.text)]
        except (ValueError, IndexError):
            return Opaque(f"missing shared string {v.text}")
        return Text(s) if s.startswith("=") or s in ERRORS else s
    if t == "str":
        return Text(v.text) if v.text.startswith("=") or v.text in ERRORS else v.text
    if t == "b":
        return v.text.strip() in ("1", "true")
    if t == "e":
        return v.text
    if t == "d":
        try:
            return from_ISO8601(v.text)
        except ValueError:
            return Opaque(f"date {v.text!r}")
    n = _number(v.text)
    return DateNum(n) if c.get("s") in dates and isinstance(n, (int, float)) else n


def read_package_cells(pkg):
    """{sheet: {coord: value}} in workbook order."""
    sst = [_text(si) for si in pkg.shared_strings()]
    dates = _date_styles(pkg)
    return {name: sheet_cells(pkg, part, sst, dates)
            for name, _, part in pkg.sheets() if part in pkg.parts and _root_name(pkg.parts[part]) in SHEET_ROOTS}


def _root_name(data):
    """Local name of an XML document's root element, without parsing it all."""
    for _, el in etree.iterparse(io.BytesIO(data), events=("start",), resolve_entities=False, no_network=True):
        return local(el)


def read_cells(path):
    """Return {sheet: {coord: value}} where formulas stay as '=...' strings."""
    return read_package_cells(Package.open(path))


def fmt(v):
    if v is None:
        return "(empty)"
    if isinstance(v, ArrayF):
        return "{" + v.text + "}"
    if isinstance(v, Opaque):
        return f"<{v.desc}>"
    if isinstance(v, Text):
        return repr(str(v))
    if isinstance(v, DateNum):
        return v.display()
    return repr(v) if isinstance(v, str) and not v.startswith("=") else str(v)


def _cell_sort_key(coord):
    col, row = coordinate_from_string(coord)
    return (row, column_index_from_string(col))


# ---------- normalizing objects for comparison ----------

def _strip_chart_caches(root):
    """Charts embed cached copies of the cells they plot; drop them so a
    cell edit doesn't look like a chart edit."""
    for el in list(root.iter(f"{{{CHART}}}numCache", f"{{{CHART}}}strCache")):
        el.getparent().remove(el)
    # Axis ids only link a chart's axes to each other, and some writers derive
    # them from the chart's file number; compare them by order of appearance.
    ids = {}
    for el in root.iter(f"{{{CHART}}}axId", f"{{{CHART}}}crossAx"):
        el.set("val", str(ids.setdefault(el.get("val"), len(ids))))


def _norm_cache(root, derived):
    for a in VOLATILE_CACHE_ATTRS:
        root.attrib.pop(a, None)
    if derived:  # the shared items are just the distinct values of the source data
        for si in root.iter(m("sharedItems")):
            si.clear()


def _norm_pivot(root, derived):
    for a in ("cacheId", "updatedVersion"):
        root.attrib.pop(a, None)
    if derived:  # what a refresh rewrites: rendered rows/columns, size, item lists
        for tag in ("rowItems", "colItems"):
            for el in root.findall(m(tag)):
                root.remove(el)
        loc = root.find(m("location"))
        if loc is not None:
            loc.attrib.pop("ref", None)
        for items in root.iter(m("items")):
            hidden = sum(1 for i in items if i.get("h") in ("1", "true"))
            items.clear()
            items.set("hidden", str(hidden))


def normalized(keys, part, derived=False):
    """Comparable form of a part: relationship ids replaced by what they point
    at, and noise (chart caches, refresh stamps) removed. derived=True also
    drops everything a pivot refresh recomputes from the source data."""
    if part is None:
        return None
    pkg = keys.pkg
    rels = {rid: (typ, t, ext) for rid, typ, t, ext in pkg.rels(part)}
    target_key = lambda t, ext: t if ext else repr(keys.key_of.get(t, t))
    relset = tuple(sorted((kind_of(typ), target_key(t, ext)) for typ, t, ext in rels.values()
                          if not (derived and kind_of(typ) == "pivotCacheRecords")))
    try:
        root = parse_xml(pkg.parts[part])
    except etree.XMLSyntaxError:
        return pkg.parts[part], relset
    kind = keys.key_of.get(part, ("",))[0]
    if kind == "cache":
        _norm_cache(root, derived)
    elif kind == "pivot":
        _norm_pivot(root, derived)
    elif local(root) == "chartSpace":
        _strip_chart_caches(root)
    for el in root.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith(f"{{{REL}}}") and v in rels:
                el.set(k, target_key(rels[v][1], rels[v][2]))
    return etree.tostring(root), relset


# ---------- describing charts, images, comments, tables, pivots ----------

def _chart_summary(pkg, name):
    root = pkg.xml(name)
    plot = root.find(f".//{{{CHART}}}plotArea")
    kinds = [local(e) for e in plot if local(e).endswith("Chart")] if plot is not None else []
    title_el = root.find(f"{{{CHART}}}chart/{{{CHART}}}title")
    title = "".join(t.text or "" for t in title_el.iter(f"{{{DRAW}}}t")) if title_el is not None else ""
    refs = [f.text for f in root.iter(f"{{{CHART}}}f") if f.text]
    _strip_chart_caches(root)
    return f"chart: {'+'.join(kinds) or 'chart'} {title!r} data={','.join(refs)} #{digest(etree.tostring(root))}"


def _pivot_summary(pkg, name):
    p = pkg.xml(name)
    cache = pkg.target(name, "pivotCacheDefinition")
    fields = [cf.get("name") for cf in pkg.xml(cache).iter(m("cacheField"))] if cache in pkg.parts else []
    field = lambda x: fields[int(x)] if x and x.isdigit() and int(x) < len(fields) else "Values"
    axis = lambda tag: [field(f.get("x")) for el in p.findall(m(tag)) for f in el]
    values = [d.get("name") for d in p.iter(m("dataField"))]
    loc = p.find(m("location"))
    where = loc.get("ref") if loc is not None else "?"
    _norm_pivot(p, derived=True)
    return (f"pivot table: {p.get('name')} at {where} "
            f"rows={axis('rowFields')} cols={axis('colFields')} values={values} #{digest(etree.tostring(p))}")


def describe_objects(pkg):
    """{sheet: [one line per chart / image / comment / table / pivot table]}."""
    out = {}
    for sheet, _, part in pkg.sheets():
        items = []
        for _, typ, tgt, ext in pkg.rels(part):
            if ext or tgt not in pkg.parts:
                continue
            kind = kind_of(typ)
            if kind == "drawing":
                for _, t2, tgt2, ext2 in pkg.rels(tgt):
                    if ext2 or tgt2 not in pkg.parts:
                        continue
                    if kind_of(t2) == "chart":
                        items.append(_chart_summary(pkg, tgt2))
                    elif kind_of(t2) == "image":
                        items.append(f"image: {posixpath.basename(tgt2)} #{digest(pkg.parts[tgt2])}")
            elif kind == "comments":
                for c in pkg.xml(tgt).iter(m("comment")):
                    text = "".join(t.text or "" for t in c.iter(m("t")))
                    items.append(f"comment {c.get('ref')}: {text!r}")
            elif kind == "table":
                t = pkg.xml(tgt)
                cols = [c.get("name") for c in t.iter(m("tableColumn"))]
                items.append(f"table: {t.get('displayName')} {t.get('ref')} columns={cols}")
            elif kind == "pivotTable":
                items.append(_pivot_summary(pkg, tgt))
        out[sheet] = items
    return out


# ---------- textconv ----------

def textconv(path):
    objects = describe_objects(Package.open(path))
    for sheet, cells in read_cells(path).items():
        print(f"=== sheet: {sheet} ===")
        for coord, v in cells.items():
            print(f"{sheet}!{coord}\t{fmt(v)}")
        for item in objects.get(sheet, []):
            print(f"{sheet} [{item}]")


# ---------- diff ----------

def diff_cells(old, new):
    """Yield (kind, sheet, coord, old_value, new_value). Inserted, deleted and
    moved rows are reported as rows, not as every cell below them changing."""
    for sheet in sorted(set(old) | set(new)):
        if sheet not in old:
            yield ("sheet added", sheet, "", None, None)
        elif sheet not in new:
            yield ("sheet removed", sheet, "", None, None)
            continue
        yield from ((k, sheet, c, a, b) for k, c, a, b in diff_sheet(old.get(sheet, {}), new.get(sheet, {})))


def diff_sheet(o, n, positions=False):
    """[(kind, coord, old, new)] for one sheet's cells. With positions, each
    comes as (row in the new sheet, change); a deleted row sits between rows."""
    plain = [(coordinate_from_string(coord)[1], ("added" if ov is None else "removed" if nv is None else "changed",
                                                 coord, ov, nv))
             for coord in sorted(set(o) | set(n), key=_cell_sort_key)
             for ov, nv in [(o.get(coord), n.get(coord))] if not same(ov, nv)]
    out = plain
    if plain:
        aligned = _aligned_diff(_by_row(o), _by_row(n))
        # Rows only help when they explain the change more simply.
        if aligned is not None and len(aligned) < len(plain):
            out = aligned
    return out if positions else [c for _, c in out]


ROW_EVENTS = ("row inserted", "row deleted", "row moved", "rows inserted", "rows deleted")


def _by_row(cells):
    rows = {}
    for coord, v in cells.items():
        col, row = coordinate_from_string(coord)
        rows.setdefault(row, {})[column_index_from_string(col)] = v
    return rows


_ANCHOR = 500000  # a row far from any edge, so shifted references stay valid


def _shape(v, row, cache):
    """The value as it would read from any row: a formula's relative
    references are rewritten as if its cell sat on row _ANCHOR, so =B6*C6 in
    row 6 and =B7*C7 in row 7 have the same shape."""
    if type(v) is bool:
        return ("b", v)
    if isinstance(v, (int, float)):
        return ("n", f"{v:.15g}")
    if isinstance(v, Text):
        return ("t", str(v))
    if isinstance(v, ArrayF) or (isinstance(v, str) and v.startswith("=") and len(v) > 1):
        text = v.text if isinstance(v, ArrayF) else v
        key = (text, row)
        if key not in cache:
            try:
                cache[key] = Translator(text, f"A{row}").translate_formula(f"A{_ANCHOR}")
            except Exception:
                cache[key] = text
        return ("f", cache[key])
    return ("v", v)


def _align(orows, nrows, cache):
    """Line up two sheets' rows by content. Returns sorted row numbers of
    each side and index lists: pairs (same row), lone_old (deleted),
    lone_new (inserted), moved (i, j)."""
    import difflib
    shape_row = lambda row, r: tuple(sorted((c, _shape(v, r, cache)) for c, v in row.items()))
    ro, rn = sorted(orows), sorted(nrows)
    ok = [shape_row(orows[r], r) for r in ro]
    nk = [shape_row(nrows[r], r) for r in rn]
    pairs, lone_old, lone_new = [], [], []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ok, nk, autojunk=False).get_opcodes():
        if tag == "equal" or (tag == "replace" and i2 - i1 == j2 - j1):
            pairs += zip(range(i1, i2), range(j1, j2))
            continue
        # Uneven block: pair rows that still share most of their cells.
        j = j1
        for i in range(i1, i2):
            best = None
            for jj in range(j, j2):
                common = len(set(ok[i]) & set(nk[jj]))
                if common * 2 >= max(len(ok[i]), len(nk[jj]), 1):
                    best = jj
                    break
            if best is None:
                lone_old.append(i)
            else:
                lone_new += range(j, best)
                pairs.append((i, best))
                j = best + 1
        lone_new += range(j, j2)
    # A deleted row that reappears unchanged elsewhere was moved.
    moved = []
    for i in list(lone_old):
        match = next((j for j in lone_new if nk[j] == ok[i] and ok[i]), None)
        if match is not None:
            lone_old.remove(i)
            lone_new.remove(match)
            moved.append((i, match))
    return ro, rn, pairs, lone_old, lone_new, moved


def _aligned_diff(orows, nrows, limit=50000):
    if len(orows) > limit or len(nrows) > limit:
        return None
    cache = {}
    ro, rn, pairs, lone_old, lone_new, moved = _align(orows, nrows, cache)
    def summary(row, show=6):
        cols = sorted(row)
        text = ", ".join(f"{get_column_letter(c)}: {fmt(row[c])}" for c in cols[:show])
        return text + (f" (+{len(cols) - show} more)" if len(cols) > show else "")
    events = []  # (new row position, order, change)
    for i, j in pairs:
        r_o, r_n = ro[i], rn[j]
        a, b = orows[r_o], nrows[r_n]
        for c in sorted(set(a) | set(b)):
            av, bv = a.get(c), b.get(c)
            if r_o == r_n:
                if same(av, bv):
                    continue
            elif (av is None) == (bv is None) and (av is None or _shape(av, r_o, cache) == _shape(bv, r_n, cache)):
                continue
            kind = "added" if av is None else "removed" if bv is None else "changed"
            events.append((r_n, c, (kind, f"{get_column_letter(c)}{r_n}", av, bv)))
    for i in lone_old:
        # Place a deleted row where it would have been in the new sheet.
        after = [rn[j] for i2, j in pairs if i2 < i]
        events.append(((after[-1] if after else 0) + 0.5, 0, ("row deleted", f"row {ro[i]}", summary(orows[ro[i]]), None)))
    for j in lone_new:
        events.append((rn[j], 0, ("row inserted", f"row {rn[j]}", None, summary(nrows[rn[j]]))))
    for i, j in moved:
        events.append((rn[j], 0, ("row moved", f"row {ro[i]} -> {rn[j]}", summary(orows[ro[i]]), None)))
    # Empty rows inserted or deleted show up as a jump in the row offset.
    content_ins = {rn[j] for j in lone_new} | {rn[j] for _, j in moved}
    content_del = {ro[i] for i in lone_old} | {ro[i] for i, _ in moved}
    ro0, rn0 = [0] + ro, [0] + rn  # a virtual row 0 lines up the tops
    anchored = [(0, 0)] + [(i + 1, j + 1) for i, j in pairs]
    for (i1, j1), (i2, j2) in zip(anchored, anchored[1:]):
        o1, o2, n1, n2 = ro0[i1], ro0[i2], rn0[j1], rn0[j2]
        blank = (n2 - n1) - (o2 - o1) - sum(1 for r in content_ins if n1 < r < n2) \
            + sum(1 for r in content_del if o1 < r < o2)
        if blank > 0:
            events.append((n2 - 0.5, 0, ("rows inserted", f"above row {n2}", None, f"{blank} empty row(s)")))
        elif blank < 0:
            events.append((n2 - 0.5, 0, ("rows deleted", f"above row {n2}", f"{-blank} empty row(s)", None)))
    events.sort(key=lambda e: (e[0], e[1]))
    return [(e[0], e[2]) for e in events]


def diff_objects(old, new):
    for sheet in sorted(set(old) | set(new)):
        o, n = list(old.get(sheet, [])), list(new.get(sheet, []))
        for item in list(o):
            if item in n:
                n.remove(item)
                o.remove(item)
        # A removed and an added object that are the same thing (the comment
        # on A1, table Sales, the only chart) is one changed object.
        for key in dict.fromkeys(_object_key(i) for i in o):
            olds = [i for i in o if _object_key(i) == key]
            news = [i for i in n if _object_key(i) == key]
            if len(olds) == len(news) == 1:
                yield ("object changed", sheet, "", olds[0], news[0])
                o.remove(olds[0])
                n.remove(news[0])
        for item in o:
            yield ("object removed", sheet, "", item, None)
        for item in n:
            yield ("object added", sheet, "", None, item)


def _object_key(item):
    kind, _, rest = item.partition(":")
    if kind.startswith("comment") or kind in ("chart", "image"):
        return kind
    return f"{kind}: {rest.split()[0] if rest.split() else ''}"  # table / pivot table by name


def diff(old_path, new_path, markdown=False, title=None):
    old, new = Package.open(old_path), Package.open(new_path)
    # A renamed sheet keeps its sheetId: compare it under its old name and
    # report the rename, not a removed sheet plus an added one.
    renamed = {a: b for a, b in _pair(old.sheets(), new.sheets()).items() if a != b}
    back = {b: a for a, b in renamed.items()}
    as_old = lambda d: {back.get(k, k): v for k, v in d.items()}
    changes = [("sheet renamed", a, "", a, b) for a, b in renamed.items()]
    changes += list(diff_cells(read_package_cells(old), as_old(read_package_cells(new))))
    changes += list(diff_objects(describe_objects(old), as_old(describe_objects(new))))
    changes = [(k, renamed.get(s, s) if k != "sheet renamed" else s, c, o, n,
                k.startswith(("object", "sheet renamed")) or k in ROW_EVENTS)
               for k, s, c, o, n in changes]
    show = lambda v, raw: "" if v is None else v if raw else fmt(v)
    if markdown:
        print(f"### {title or new_path}\n")
        if not changes:
            print("No cell or object changes (formatting may still differ).\n")
            return 0
        print(f"{len(changes)} change(s)\n")
        print("| Sheet | Cell | Change | Before | After |")
        print("|---|---|---|---|---|")
        for kind, sheet, coord, ov, nv, raw in changes[:500]:
            esc = lambda v: str(v).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")
            print(f"| {esc(sheet)} | {coord} | {kind} | {esc(show(ov, raw))} | {esc(show(nv, raw))} |")
        if len(changes) > 500:
            print(f"\n...and {len(changes) - 500} more.")
        print()
    else:
        if not changes:
            print("No cell or object changes (formatting may still differ).")
        for kind, sheet, coord, ov, nv, raw in changes:
            if kind == "sheet renamed":
                print(f"{kind:14} {ov} -> {nv}")
            elif kind in ROW_EVENTS:
                print(f"{kind:14} {sheet} {coord}  {nv if nv is not None else ov}")
            elif raw:  # a chart, table, comment...
                what = f"{ov} -> {nv}" if ov is not None and nv is not None else nv if nv is not None else ov
                print(f"{kind:14} {sheet}  {what}")
            elif not coord:  # a whole sheet
                print(f"{kind:14} {sheet}")
            else:
                print(f"{kind:14} {sheet}!{coord}  {fmt(ov)} -> {fmt(nv)}")
    return 1 if changes else 0


# ---------- following rows through a merge ----------

class RowMap:
    """Where each row of the base version ended up on a branch that inserted,
    deleted or moved rows: r -> new row number, or None if it was deleted."""

    def __init__(self, pairs, deleted):
        self.exact = dict(pairs)
        self.deleted = set(deleted)
        self.anchors = sorted(self.exact.items())
        self._keys = [b for b, _ in self.anchors]

    def __call__(self, r):
        if r in self.exact:
            return self.exact[r]
        if r in self.deleted:
            return None
        # An empty row moves with the row above it.
        import bisect
        i = bisect.bisect_right(self._keys, r) - 1
        if i < 0:
            return r
        b, o = self.anchors[i]
        return r + (o - b)

    def start(self, r):
        """First surviving row at or after r (a range's top edge)."""
        for x in range(r, r + 100000):
            if x not in self.deleted:
                return self(x)
        return None

    def end(self, r):
        """Last surviving row at or before r (a range's bottom edge)."""
        for x in range(r, max(r - 100000, 0), -1):
            if x not in self.deleted:
                return self(x)
        return None


def row_map(base_cells, cells):
    """A RowMap if this version of a sheet inserted, deleted or moved rows
    relative to base, else None (only cells changed)."""
    if not base_cells or not cells or same_cells(base_cells, cells):
        return None
    if not any(k in ROW_EVENTS for k, *_ in diff_sheet(base_cells, cells)):
        return None
    brows, rows = _by_row(base_cells), _by_row(cells)
    cache = {}
    ro, rn, pairs, lone_old, lone_new, moved = _align(brows, rows, cache)
    # Moved rows count as deleted and inserted: references to them can't be
    # followed the way Excel would, so edits there conflict instead.
    deleted = [ro[i] for i in lone_old] + [ro[i] for i, _ in moved]
    rmap = RowMap([(ro[i], rn[j]) for i, j in pairs], deleted)
    # The alignment is a guess that reads well in a diff; a merge that follows
    # a wrong guess moves real data. Only trust it when it lines up clearly
    # more cells with base than leaving every row where it was. A real insert
    # shifts everything below it and wins easily; a handful of edits on rows
    # that look alike (same formulas down a column) can't.
    gain = _row_matches(brows, rows, rmap, cache) - _row_matches(brows, rows, lambda r: r, cache)
    if gain <= 0:
        return None
    return rmap


def _row_matches(brows, rows, where, cache):
    """How many base cells hold the same thing (formulas by shape) at the row
    `where` puts them."""
    n = 0
    for r, cells in brows.items():
        to = where(r)
        target = rows.get(to) if to is not None else None
        if not target:
            continue
        for c, v in cells.items():
            if c in target and _shape(v, r, cache) == _shape(target[c], to, cache):
                n += 1
    return n


_REF_CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_REF_ROW = re.compile(r"^(\$?)(\d+)$")


def _move_ref(part, fn):
    for rx, row_group in ((_REF_CELL, 4), (_REF_ROW, 2)):
        mt = rx.match(part)
        if mt:
            r = fn(int(mt.group(row_group)))
            if r is None:
                return None
            groups = list(mt.groups())
            groups[row_group - 1] = str(r)
            return "".join(groups)
    return part  # a whole column, a name: no rows to move


def _ref_row(part):
    mt = _REF_CELL.match(part) or _REF_ROW.match(part)
    return int(mt.groups()[-1]) if mt else 0


def rewrite_refs(formula, host, maps):
    """Renumber a formula's row references for rows that moved, as Excel does
    when rows are inserted or deleted. maps: {sheet name: RowMap}; host is
    the sheet the formula sits on."""
    if not maps or not isinstance(formula, str) or isinstance(formula, Text) or not formula.startswith("="):
        return formula
    try:
        from openpyxl.formula.tokenizer import Tokenizer
        tok = Tokenizer(formula)
    except Exception:
        return formula
    changed = False
    for t in tok.items:
        if t.type != "OPERAND" or t.subtype != "RANGE":
            continue
        prefix, ref = (t.value.rsplit("!", 1) if "!" in t.value else ("", t.value))
        if "[" in prefix or "[" in ref:
            continue  # another workbook, or a table reference
        sheet = prefix[1:-1].replace("''", "'") if prefix.startswith("'") else (prefix or host)
        m = maps.get(sheet)
        if m is None:
            continue
        parts = ref.split(":")
        if len(parts) == 1:
            new = _move_ref(parts[0], m)
        elif len(parts) == 2:
            a, b = _move_ref(parts[0], m.start), _move_ref(parts[1], m.end)
            new = None if a is None or b is None or _ref_row(a) > _ref_row(b) else f"{a}:{b}"
        else:
            continue
        value = (prefix + "!" if prefix else "") + (new if new is not None else "#REF!")
        if new is None:
            value = "#REF!"
        if value != t.value:
            t.value = value
            changed = True
    return tok.render() if changed else formula


# ---------- visual diff ----------

_CSS = """
:root{--bg:#fff;--fg:#1f2328;--muted:#656d76;--line:#d0d7de;--head:#f6f8fa;
--chg:#fff1c2;--chg-b:#d4a72c;--add:#dafbe1;--add-b:#2da44e;--del:#ffebe9;--del-b:#cf222e;--mov:#ddf4ff;--mov-b:#0969da}
@media (prefers-color-scheme:dark){:root{--bg:#0d1117;--fg:#e6edf3;--muted:#8d96a0;--line:#30363d;--head:#161b22;
--chg:#3b2e0a;--chg-b:#d29922;--add:#0f2e1a;--add-b:#3fb950;--del:#3a1418;--del-b:#f85149;--mov:#0c2d4a;--mov-b:#4493f8}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.45 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1400px;margin:0 auto;padding:24px 16px 64px}h1{font-size:20px;margin:0 0 4px}h2{font-size:17px;margin:32px 0 8px}
h3{font-size:15px;margin:20px 0 8px}.muted{color:var(--muted)}.legend span{display:inline-block;margin-right:14px}
.sw{display:inline-block;width:12px;height:12px;border-radius:3px;vertical-align:-1px;margin-right:5px;border:1px solid}
.wrap{width:fit-content;max-width:100%;overflow-x:auto;border:1px solid var(--line);border-radius:8px}table{border-collapse:collapse;font-size:13px}
th,td{border:1px solid var(--line);padding:3px 8px;white-space:nowrap;max-width:260px;overflow:hidden;text-overflow:ellipsis}
th{background:var(--head);color:var(--muted);font-weight:500;position:sticky;top:0}td.rn{background:var(--head);color:var(--muted);text-align:right}
td.f{font-family:ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;font-size:12px}
td.chg{background:var(--chg);box-shadow:inset 0 0 0 1px var(--chg-b)}td.add{background:var(--add);box-shadow:inset 0 0 0 1px var(--add-b)}
td.rem{background:var(--del);text-decoration:line-through;box-shadow:inset 0 0 0 1px var(--del-b)}
tr.ins td{background:var(--add)}tr.dele td{background:var(--del);text-decoration:line-through}tr.mov td{background:var(--mov)}
tr.gap td{background:var(--bg);color:var(--muted);text-align:left;padding-left:48px;font-style:italic;border-left:0;border-right:0}
.was{display:block;font-size:11px;color:var(--muted);text-decoration:line-through}
ul.obj{margin:6px 0;padding-left:20px}ul.obj li{margin:2px 0}code{font-family:ui-monospace,Menlo,Consolas,monospace;font-size:12px}
a{color:var(--mov-b)}nav a{margin-right:12px}
"""


def _cell_html(v):
    """(text, is_formula) for showing a cell value."""
    if v is None:
        return "", False
    if type(v) is bool:
        return ("TRUE" if v else "FALSE"), False
    if isinstance(v, DateNum):
        return v.display(), False
    if isinstance(v, float):
        return f"{v:.15g}", False
    if isinstance(v, ArrayF):
        return "{" + v.text + "}", True
    if isinstance(v, Opaque):
        return f"<{v.desc}>", False
    if isinstance(v, str) and not isinstance(v, Text) and v.startswith("="):
        return v, True
    return str(v), False


def _sheet_html(o, n, context=2, max_rows=1500, max_cols=40):
    import html as H
    changes = diff_sheet(o, n, positions=True)
    if not changes:
        return "", 0
    orows, nrows = _by_row(o), _by_row(n)
    cells, rowcls, deleted, notes = {}, {}, [], []
    for pos, (kind, coord, ov, nv) in changes:
        if kind in ("changed", "added", "removed"):
            col, r = coordinate_from_string(coord)
            cells[(r, column_index_from_string(col))] = (kind, ov)
        elif kind == "row inserted":
            rowcls[int(coord.split()[1])] = ("ins", "inserted")
        elif kind == "row moved":
            a, b = coord.split()[1], coord.split()[3]
            rowcls[int(b)] = ("mov", f"moved from row {a}")
        elif kind == "row deleted":
            deleted.append((pos, int(coord.split()[1])))
        else:
            notes.append((pos, f"{nv or ov} {'inserted' if kind == 'rows inserted' else 'deleted'} here"))
    marked = {r for r, _ in cells} | set(rowcls)
    # Rows around each change, and around where deleted rows used to be.
    between = {int(p) for p, _ in deleted + notes} | {int(p) + 1 for p, _ in deleted + notes}
    show = {r + d for r in marked | between for d in range(-context, context + 1)
            if r + d in nrows or r + d in marked}
    show = {r for r in show if r >= 1}
    items = [(r, 1, r) for r in show] + [(p, 0, ("del", r)) for p, r in deleted] + [(p, 0, ("note", t)) for p, t in notes]
    items.sort(key=lambda x: (x[0], x[1]))
    truncated = len(items) > max_rows
    items = items[:max_rows]
    used = set()
    for _, _, it in items:
        row = nrows.get(it) if isinstance(it, int) else orows.get(it[1]) if it[0] == "del" else None
        used |= set(row or ())
    used |= {c for _, c in cells}
    if not used:
        used = {1}
    cols = list(range(1, max(used) + 1))
    if len(cols) > max_cols:  # wide sheet: changed columns and their neighbours
        keep = {c + d for c in {c for _, c in cells} | {1} for d in (-1, 0, 1)}
        cols = [c for c in cols if c in keep][:max_cols]
    out = ['<div class="wrap"><table><thead><tr><th></th>']
    prev = None
    for c in cols:
        if prev is not None and c != prev + 1:
            out.append('<th>…</th>')
        out.append(f"<th>{get_column_letter(c)}</th>")
        prev = c
    out.append("</tr></thead><tbody>")
    span = len(cols) + 1 + sum(1 for a, b in zip(cols, cols[1:]) if b != a + 1)
    last = 0
    for _, _, it in items:
        if isinstance(it, int):
            if it > last + 1:
                out.append(f'<tr class="gap"><td colspan="{span}">⋯ {it - last - 1} unchanged row(s)</td></tr>')
            last = it
            cls, label = rowcls.get(it, ("", ""))
            row = nrows.get(it, {})
            out.append(f'<tr class="{cls}"><td class="rn" title="{H.escape(label)}">{it}</td>')
            prev = None
            for c in cols:
                if prev is not None and c != prev + 1:
                    out.append("<td></td>")
                prev = c
                text, is_f = _cell_html(row.get(c))
                mark = cells.get((it, c))
                klass = ["f"] if is_f else []
                was = ""
                if mark:
                    kind, ov = mark
                    klass.append({"changed": "chg", "added": "add", "removed": "rem"}[kind])
                    if kind == "removed":
                        text, _ = _cell_html(ov)
                    elif kind == "changed":
                        was = f'<span class="was">{H.escape(_cell_html(ov)[0])}</span>'
                title = f' title="was: {H.escape(_cell_html(mark[1])[0])}"' if mark and mark[0] == "changed" else ""
                out.append(f'<td class="{" ".join(klass)}"{title}>{was}{H.escape(text)}</td>')
            out.append("</tr>")
        elif it[0] == "del":
            row = orows.get(it[1], {})
            out.append(f'<tr class="dele"><td class="rn" title="deleted">−{it[1]}</td>')
            prev = None
            for c in cols:
                if prev is not None and c != prev + 1:
                    out.append("<td></td>")
                prev = c
                text, is_f = _cell_html(row.get(c))
                out.append(f'<td class="{"f" if is_f else ""}">{H.escape(text)}</td>')
            out.append("</tr>")
        else:
            out.append(f'<tr class="gap"><td colspan="{span}">{H.escape(it[1])}</td></tr>')
    out.append("</tbody></table></div>")
    if truncated:
        out.append(f'<p class="muted">Showing the first {max_rows} rows of changes.</p>')
    return "".join(out), len(changes)


def html_report(items, title="Excel changes"):
    """A self-contained page: items are (name, old Package, new Package)."""
    import html as H
    body, nav = [], []
    for k, (name, old, new) in enumerate(items):
        renamed = {a: b for a, b in _pair(old.sheets(), new.sheets()).items() if a != b}
        back = {b: a for a, b in renamed.items()}
        oc = read_package_cells(old)
        nc = {back.get(s, s): v for s, v in read_package_cells(new).items()}
        objs = list(diff_objects(describe_objects(old), {back.get(s, s): v for s, v in describe_objects(new).items()}))
        order = [back.get(n_, n_) for n_, _, _ in new.sheets()] + [n_ for n_, _, _ in old.sheets() if n_ not in nc]
        parts, total = [], 0
        for sheet in dict.fromkeys(order):
            shown = renamed.get(sheet, sheet)
            head = H.escape(shown)
            if sheet in renamed:
                head += f' <span class="muted">(renamed from {H.escape(sheet)})</span>'
                total += 1
            if sheet not in oc and sheet in nc:
                head += ' <span class="muted">(new sheet)</span>'
                total += 1
            if sheet in oc and sheet not in nc:
                parts.append(f"<h3>{head} <span class=\"muted\">(sheet deleted)</span></h3>")
                total += 1
                continue
            grid, n_ = _sheet_html(oc.get(sheet, {}), nc.get(sheet, {}))
            sheet_objs = [o_ for o_ in objs if o_[1] == sheet]
            if not grid and not sheet_objs and sheet not in renamed and sheet in oc:
                continue
            total += n_ + len(sheet_objs)
            parts.append(f"<h3>{head}</h3>{grid}")
            if sheet_objs:
                parts.append('<ul class="obj">' + "".join(
                    f"<li>{H.escape(kind)}: <code>{H.escape(str(nv if nv is not None else ov))}</code>"
                    + (f' <span class="muted">(was <code>{H.escape(str(ov))}</code>)</span>' if ov is not None and nv is not None else "")
                    + "</li>" for kind, _, _, ov, nv in sheet_objs) + "</ul>")
        anchor = f"wb{k}"
        nav.append(f'<a href="#{anchor}">{H.escape(name)}</a>')
        body.append(f'<h2 id="{anchor}">{H.escape(name)} <span class="muted">· {total} change(s)</span></h2>')
        body.append("".join(parts) or '<p class="muted">No cell or object changes (formatting may still differ).</p>')
    legend = ('<p class="legend muted"><span><i class="sw" style="background:var(--chg);border-color:var(--chg-b)"></i>changed (old value above)</span>'
              '<span><i class="sw" style="background:var(--add);border-color:var(--add-b)"></i>added / inserted row</span>'
              '<span><i class="sw" style="background:var(--del);border-color:var(--del-b)"></i>removed / deleted row</span>'
              '<span><i class="sw" style="background:var(--mov);border-color:var(--mov-b)"></i>moved row</span></p>')
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f"<title>{H.escape(title)}</title><style>{_CSS}</style></head><body><main>"
            f"<h1>{H.escape(title)}</h1>{legend}"
            + (f"<nav>{''.join(nav)}</nav>" if len(nav) > 1 else "")
            + "".join(body) + f'<p class="muted">Made by xlgit {__version__}.</p></main></body></html>')


def write_html(items, out=None, open_it=True, title="Excel changes"):
    import tempfile
    import webbrowser
    page = html_report(items, title)
    if out is None:
        fd, out = tempfile.mkstemp(prefix="xlgit-diff-", suffix=".html")
        os.close(fd)
    with open(out, "w", encoding="utf-8") as f:
        f.write(page)
    print(f"Visual diff: {os.path.abspath(out)}")
    if open_it:
        try:
            webbrowser.open("file://" + os.path.abspath(out).replace("\\", "/"))
        except Exception:
            pass
    return out


# ---------- merge ----------

def _pair(a, b):
    """Map sheet names in a to names in b: same name first, then same sheetId
    (Excel keeps sheetId when a sheet is renamed)."""
    names_b = {n for n, _, _ in b}
    out = {n: n for n, _, _ in a if n in names_b}
    used = set(out.values())
    by_id = {sid: n for n, sid, _ in b if n not in used}
    for n, sid, _ in a:
        if n not in out and sid in by_id and by_id[sid] not in used:
            out[n] = by_id[sid]
            used.add(by_id[sid])
    return out


def pivot_locations(pkg, sheet_part):
    refs = []
    for _, typ, t, ext in (pkg.rels(sheet_part) if sheet_part else []):
        if not ext and kind_of(typ) == "pivotTable" and t in pkg.parts:
            loc = pkg.xml(t).find(m("location"))
            if loc is not None and loc.get("ref"):
                refs.append(loc.get("ref"))
    return refs


def _three_way(b, o, t):
    """(value, ok): the merged value, or ok=False if both sides changed it differently."""
    if t == b or t == o:
        return o, True
    if o == b:
        return t, True
    return o, False


class Merger:
    """3-way merge of BASE/OURS/THEIRS; the result starts as a copy of OURS'
    zip and only the XML that has to change gets rewritten."""

    def __init__(self, base_path, ours_path, theirs_path):
        self.b, self.o, self.t = (Package.open(p) for p in (base_path, ours_path, theirs_path))
        self.bc, self.oc, self.tc = (read_package_cells(p) for p in (self.b, self.o, self.t))
        self.notes = []
        self.swapped = False
        self.plan_rows()
        self.res = dict(self.o.parts)
        self.order = list(self.o.order)
        self.trees = {}
        self.imported = {}  # their part name -> part name in the result
        self.conflicts = []
        self.cells_taken = 0
        self.objects_taken = []
        self.stale_pivots = False
        self.touched_tables = set()
        self.table_renames = {}
        self.wb = self.o.workbook
        self.same_styles = self.o.styles_bytes() == self.t.styles_bytes()
        self.sst = self.o.shared_strings()
        wbpr = self.o.xml(self.wb).find(m("workbookPr"))
        self.epoch = CALENDAR_MAC_1904 if wbpr is not None and wbpr.get("date1904") in ("1", "true") \
            else CALENDAR_WINDOWS_1900

    # --- result tree plumbing ---

    def tree(self, name, create=None):
        if name not in self.trees:
            if name in self.res:
                self.trees[name] = parse_xml(self.res[name])
            elif create is not None:
                self.trees[name] = create
                self.put(name, b"")
            else:
                return None
        return self.trees[name]

    def put(self, name, data):
        self.res[name] = data
        self.trees.pop(name, None) if data else None
        if name not in self.order:
            self.order.append(name)

    def drop(self, name):
        self.res.pop(name, None)
        self.trees.pop(name, None)
        if name in self.order:
            self.order.remove(name)
        ct = self.tree("[Content_Types].xml")
        for o in list(ct.iter(f"{{{CTYPES}}}Override")):
            if o.get("PartName").lstrip("/") == name:
                ct.remove(o)

    def rels_tree(self, owner):
        return self.tree(rels_name(owner), create=etree.Element(f"{{{PKG_REL}}}Relationships", nsmap={None: PKG_REL}))

    def rels_of(self, owner):
        rt = self.tree(rels_name(owner))
        if rt is None:
            return []
        out = []
        for r in rt:
            ext = r.get("TargetMode") == "External"
            out.append((r.get("Id"), r.get("Type"), r.get("Target") if ext else resolve(owner, r.get("Target")), ext))
        return out

    def add_rel(self, owner, typ, target, external=False):
        rt = self.rels_tree(owner)
        ids = {r.get("Id") for r in rt}
        n = 1
        while f"rId{n}" in ids:
            n += 1
        r = etree.SubElement(rt, f"{{{PKG_REL}}}Relationship")
        r.set("Id", f"rId{n}")
        r.set("Type", typ)
        r.set("Target", target if external else relative(target, owner))
        if external:
            r.set("TargetMode", "External")
        return f"rId{n}"

    def remove_rel(self, owner, rid):
        rt = self.rels_tree(owner)
        for r in list(rt):
            if r.get("Id") == rid:
                rt.remove(r)

    def copy_content_type(self, their_name, our_name):
        info = self.t.content_type(their_name)
        if not info:
            return
        ct = self.tree("[Content_Types].xml")
        if info[0] == "override":
            for o in ct.iter(f"{{{CTYPES}}}Override"):
                if o.get("PartName").lstrip("/") == our_name:
                    o.set("ContentType", info[1])
                    return
            o = etree.SubElement(ct, f"{{{CTYPES}}}Override")
            o.set("PartName", "/" + our_name)
            o.set("ContentType", info[1])
        elif not any(d.get("Extension").lower() == info[1] for d in ct.iter(f"{{{CTYPES}}}Default")):
            d = etree.Element(f"{{{CTYPES}}}Default")
            d.set("Extension", info[1])
            d.set("ContentType", info[2])
            ct.insert(0, d)

    def unique_name(self, name):
        d, b = posixpath.split(name)
        stem, num, ext = re.match(r"(.*?)(\d*)(\.[^.]*)?$", b).groups()
        n = int(num or 1)
        while True:
            n += 1
            cand = posixpath.join(d, f"{stem}{n}{ext or ''}")
            if cand not in self.res and cand not in self.t.parts and cand not in self.b.parts:
                return cand

    def ours_for(self, their_part):
        """Our part for the same object as one of theirs, if both branches
        inherited it from base and we still have it."""
        key = self.kt.key_of.get(their_part)
        if key is None or key not in self.kb.part_of:
            return None
        ours = self.ko.part_of.get(key)
        return ours if ours in self.res else None

    def import_part(self, their_name, keep_rel=None):
        """Copy one of their parts (and everything it links to) into the result,
        renaming anything that would collide with a part we already have."""
        if their_name in self.imported:
            return self.imported[their_name]
        ours = their_name if their_name not in self.res else self.unique_name(their_name)
        self.imported[their_name] = ours
        self.put(ours, self.t.parts[their_name])
        self.copy_content_type(their_name, ours)
        self.copy_rels(their_name, ours, keep_rel)
        return ours

    def copy_rels(self, their_owner, our_owner, keep_rel=None):
        """Their relationships for a part, retargeted into the result: links to
        objects we share point at our copy, anything new gets imported.
        Returns (rId, type) pairs that keep_rel rejected."""
        rn = rels_name(our_owner)
        rels = self.t.rels(their_owner)
        self.trees.pop(rn, None)
        if not rels:
            if rn in self.res:
                self.drop(rn)
            return []
        rt = etree.Element(f"{{{PKG_REL}}}Relationships", nsmap={None: PKG_REL})
        dropped = []
        for rid, typ, tgt, ext in rels:
            if keep_rel and not ext and not keep_rel(typ):
                dropped.append((rid, typ))
                continue
            r = etree.SubElement(rt, f"{{{PKG_REL}}}Relationship")
            r.set("Id", rid)
            r.set("Type", typ)
            if ext:
                r.set("Target", tgt)
                r.set("TargetMode", "External")
                continue
            ours = self.ours_for(tgt)
            if ours is None:
                ours = self.import_part(tgt) if tgt in self.t.parts else tgt
            r.set("Target", relative(ours, our_owner))
        self.trees[rn] = rt
        self.put(rn, b"")
        return dropped

    def take_theirs(self, our_part, their_part):
        self.put(our_part, self.t.parts[their_part])
        self.copy_rels(their_part, our_part)

    def changed_since_base(self, our_part, seen=None):
        """Did we edit this object, or anything hanging off it?"""
        seen = seen if seen is not None else set()
        if our_part in seen:
            return False
        seen.add(our_part)
        key = self.ko.key_of.get(our_part)
        base_part = self.kb.part_of.get(key)
        if base_part is None or normalized(self.ko, our_part) != normalized(self.kb, base_part):
            return True
        return any(self.changed_since_base(t, seen) for _, typ, t, ext in self.o.rels(our_part)
                   if not ext and t in self.o.parts and kind_of(typ) != "pivotCacheDefinition")

    def sheet_map(self):
        """{name: (sheet element, part)} for the result as it stands now."""
        wb = self.tree(self.wb)
        targets = {r.get("Id"): resolve(self.wb, r.get("Target")) for r in self.rels_tree(self.wb)}
        return {s.get("name"): (s, targets.get(s.get(RID))) for s in wb.find(m("sheets"))}

    def conflict(self, sheet, cell, base, ours, theirs):
        if self.swapped:  # the result was built on their layout; report from your side
            ours, theirs = theirs, ours
            flip = lambda v: v.replace("kept ours", "kept theirs (it has their row changes)") \
                if isinstance(v, str) and not isinstance(v, Text) else v
            ours, theirs = flip(ours), flip(theirs)
        self.conflicts.append((sheet, cell, base, ours, theirs))

    # --- rows that moved ---

    def plan_rows(self):
        """Find sheets where one branch inserted, deleted or moved rows and the
        other only edited cells. The merge follows the branch that moved rows:
        the other side's edits land on the rows where their cells now are,
        with formula references renumbered as Excel would. If only their
        branch moved rows, the result is built on their version (Excel already
        moved everything else consistently there) and your edits move onto it."""
        bs, os_, ts = self.b.sheets(), self.o.sheets(), self.t.sheets()
        bo, bt = _pair(bs, os_), _pair(bs, ts)
        omaps, tmaps = {}, {}
        for bname, _, _ in bs:
            b = self.bc.get(bname)
            om = row_map(b, self.oc.get(bo.get(bname))) if bo.get(bname) else None
            tm = row_map(b, self.tc.get(bt.get(bname))) if bt.get(bname) else None
            if om and tm:
                self.notes.append(f"both branches inserted or deleted rows on {bname!r}; merged cell by cell")
            elif om:
                omaps[bname] = om
            elif tm:
                tmaps[bname] = tm
        if tmaps and not omaps:
            self.swapped = True
            self.o, self.t = self.t, self.o
            self.oc, self.tc = self.tc, self.oc
            omaps, tmaps, bo, bt = tmaps, {}, bt, bo
            self.notes.append("their branch inserted or deleted rows on " + ", ".join(map(repr, omaps))
                              + "; your edits moved to the rows where those cells are now")
        for bname in tmaps:
            self.notes.append(f"their branch inserted or deleted rows on {bname!r} and yours did on other "
                              f"sheets; {bname!r} merged cell by cell")
        self.maps_b = omaps  # by base sheet name
        self.maps_t = {bt[n]: mp for n, mp in omaps.items() if bt.get(n)}  # by their sheet names

    # --- the merge ---

    def run(self):
        bs, os_, ts = self.b.sheets(), self.o.sheets(), self.t.sheets()
        bo, bt = _pair(bs, os_), _pair(bs, ts)
        self.lineage_o = {v: k for k, v in bo.items()}
        self.lineage_t = {v: k for k, v in bt.items()}
        self.kb = Keys(self.b, {n: n for n, _, _ in bs})
        self.ko = Keys(self.o, {n: self.lineage_o.get(n, "ours:" + n) for n, _, _ in os_}, self.kb.part_of)
        self.kt = Keys(self.t, {n: self.lineage_t.get(n, "theirs:" + n) for n, _, _ in ts}, self.kb.part_of)
        self.sheet_el = {n: el for n, (el, _) in self.sheet_map().items()}  # ours, before any change
        self.new_sheet_el = {}
        self.capture_names()
        self.plan_table_renames()

        done_t = set()
        for bname, _, bpart in bs:
            o, t = bo.get(bname), bt.get(bname)
            done_t.add(t)
            if o and t:
                self.merge_sheet(bname, o, t)
                self.merge_sheet_name(bname, o, t)
            elif o:
                if same_cells(self.oc.get(o), self.bc.get(bname)):
                    self.delete_sheet(o)
                else:
                    self.conflict(o, "(sheet)", "exists", "edited", "deleted")
            elif t and not same_cells(self.tc.get(t), self.bc.get(bname)):
                self.conflict(t, "(sheet)", "exists", "deleted", "edited")
        ours_from_base = set(bo.values())
        for tname, _, tpart in ts:
            if tname in done_t:
                continue
            if tname in self.sheet_el and tname not in ours_from_base:
                self.merge_sheet(None, tname, tname)  # both branches added a sheet with this name
            elif tname in self.sheet_map():
                self.conflict(tname, "(sheet)", "", "exists", "added with the same name")
            else:
                self.import_sheet(tname, tpart)

        self.merge_names()
        self.merge_object_parts()
        self.merge_tables()
        self.merge_pivots()
        self.fix_tables()
        self.fix_pivots()
        self.report_unmerged_additions()
        self.write_conflict_sheet()
        return self.finish()

    # --- cells ---

    def merge_sheet(self, bname, oname, tname):
        bpart = dict((n, p) for n, _, p in self.b.sheets()).get(bname)
        opart = dict((n, p) for n, _, p in self.o.sheets()).get(oname)
        tpart = dict((n, p) for n, _, p in self.t.sheets()).get(tname)
        # Cells a pivot table displays are regenerated when it refreshes, so
        # they never conflict: we keep ours and Excel refreshes on open.
        pivot_areas = [ref for pkg, part in ((self.b, bpart), (self.o, opart), (self.t, tpart))
                       for ref in pivot_locations(pkg, part)]
        b = self.bc.get(bname, {}) if bname else {}
        o, t = self.oc.get(oname, {}), self.tc.get(tname, {})
        rows = self.maps_b.get(bname) if bname else None
        edits, src = {}, {}
        for coord in set(b) | set(t):
            bv, tv = b.get(coord), t.get(coord)
            if same(tv, bv):
                continue
            where = coord
            if rows:
                col, r = coordinate_from_string(coord)
                if rows(r) is None:
                    self.conflict(oname, f"{coord} (row deleted)", bv, "row deleted", tv)
                    continue
                where = f"{col}{rows(r)}"
            # Compare as if both sides had the row changes.
            bv = rewrite_refs(bv, bname, self.maps_b)
            tv = rewrite_refs(tv, tname, self.maps_t)
            ov = o.get(where)
            if same(tv, ov):
                continue
            if same(ov, bv) and not isinstance(tv, Opaque):
                edits[where], src[where] = tv, coord
            elif any(overlaps(where, ref) for ref in pivot_areas):
                self.stale_pivots = True
            else:
                if self.swapped and not isinstance(tv, Opaque):
                    edits[where], src[where] = tv, coord  # your value wins a clash, as always
                self.conflict(oname, where, bv, ov, tv)
        if edits:
            self.edit_cells(oname, tname, edits, src)
        for key in ("drawing", "legacyDrawing", "comments"):
            self.merge_sheet_object(oname, key, bpart, opart, tpart)
        for kind in ("table", "pivotTable"):
            self.merge_sheet_collection(oname, kind, bpart, opart, tpart)

    def edit_cells(self, oname, tname, edits, src=None):
        """Write edits ({cell: value}) into our sheet. src maps each cell to
        where it was in their version, if rows moved."""
        src = src or {}
        at = lambda coord: src.get(coord, coord)
        part = self.sheet_map()[oname][1]
        root = self.tree(part)
        sd = root.find(m("sheetData"))
        if sd is None:
            sd = etree.SubElement(root, m("sheetData"))
            self.place(root, sd, WS_ORDER)
        self.spell_out_refs(sd)
        self.unshare_formulas(root, self.oc.get(oname, {}))
        # Style numbers index into styles.xml, so theirs only carry over when
        # both branches have the same styles.xml, and only if we didn't restyle.
        their_styles, base_styles = {}, None
        if self.same_styles:
            tpart = dict((n, p) for n, _, p in self.t.sheets())[tname]
            their_styles = self.t.scan_sheet(tpart, styles_for={at(c) for c in edits})[2]
            bpart = dict((n, p) for n, _, p in self.b.sheets()).get(self.lineage_o.get(oname))
            if bpart and self.b.styles_bytes() == self.o.styles_bytes():
                base_styles = self.b.scan_sheet(bpart, styles_for={at(c) for c in edits})[2]
        rows = {}
        last = 0
        for row in sd.findall(m("row")):
            last = int(row.get("r") or last + 1)
            rows[last] = row
        for coord, value in sorted(edits.items(), key=lambda kv: _cell_sort_key(kv[0])):
            col, rnum = coordinate_from_string(coord)
            row = self.get_row(sd, rows, rnum)
            row.attrib.pop("spans", None)
            c = self.get_cell(row, coord, column_index_from_string(col))
            restyled = base_styles is not None and c.get("s") != base_styles.get(at(coord))
            self.write_value(c, value)
            if self.same_styles and not restyled:
                s = their_styles.get(at(coord))
                c.set("s", s) if s else c.attrib.pop("s", None)
            if value is None and c.get("s") is None and len(c) == 0:
                row.remove(c)
            self.cells_taken += 1
        for rnum, row in rows.items():
            if len(row) == 0 and set(row.attrib) <= {"r", "spans"}:
                sd.remove(row)
        self.update_dimension(root)

    @staticmethod
    def spell_out_refs(sd):
        """Row and cell references are optional in the file format (the next
        row, the next column); write them all out before inserting anything."""
        rnum = 0
        for row in sd.findall(m("row")):
            rnum = int(row.get("r")) if row.get("r") else rnum + 1
            row.set("r", str(rnum))
            col = 0
            for c in row.findall(m("c")):
                if c.get("r"):
                    col = column_index_from_string(coordinate_from_string(c.get("r"))[0])
                else:
                    col += 1
                    c.set("r", f"{get_column_letter(col)}{rnum}")

    def unshare_formulas(self, root, values):
        """Shared formulas are stored once and implied for a range; editing one
        cell of the range would break the rest, so spell each one out."""
        for c in root.iter(m("c")):
            f = c.find(m("f"))
            if f is not None and f.get("t") == "shared":
                v = values.get(c.get("r"))
                if isinstance(v, str) and not isinstance(v, Text) and v.startswith("="):
                    f.text = v[1:]
                    for a in ("t", "ref", "si"):
                        f.attrib.pop(a, None)

    @staticmethod
    def get_row(sd, rows, rnum):
        row = rows.get(rnum)
        if row is None:
            row = etree.SubElement(sd, m("row"))
            row.set("r", str(rnum))
            later = [rows[n] for n in sorted(rows) if n > rnum]
            if later:
                later[0].addprevious(row)
            rows[rnum] = row
        return row

    @staticmethod
    def get_cell(row, coord, colnum):
        for c in row.findall(m("c")):
            if c.get("r") == coord:
                return c
        c = etree.SubElement(row, m("c"))
        c.set("r", coord)
        for other in row.findall(m("c")):
            if other is not c and column_index_from_string(coordinate_from_string(other.get("r"))[0]) > colnum:
                other.addprevious(c)
                break
        return c

    def write_value(self, c, v, as_text=False):
        for child in list(c):
            if local(child) in ("f", "v", "is"):
                c.remove(child)
        for a in ("t", "cm", "vm"):
            c.attrib.pop(a, None)
        if v is None:
            return
        new = []
        if isinstance(v, ArrayF):
            f = etree.SubElement(c, m("f"))
            f.set("t", "array")
            f.set("ref", v.ref)
            f.text = v.text[1:] if v.text.startswith("=") else v.text
            new.append(f)
        elif isinstance(v, bool):
            c.set("t", "b")
            new.append(self._sub(c, "v", "1" if v else "0"))
        elif isinstance(v, (int, float)):
            new.append(self._sub(c, "v", repr(int(v) if type(v) is int else float(v))))
        elif isinstance(v, (datetime.date, datetime.time, datetime.timedelta)):
            new.append(self._sub(c, "v", repr(to_excel(v, self.epoch))))
        elif v.startswith("=") and len(v) > 1 and not as_text and not isinstance(v, Text):
            new.append(self._sub(c, "f", v[1:]))
        elif v in ERRORS and not as_text and not isinstance(v, Text):
            c.set("t", "e")
            new.append(self._sub(c, "v", v))
        else:
            c.set("t", "inlineStr")
            is_ = etree.SubElement(c, m("is"))
            t = etree.SubElement(is_, m("t"))
            t.text = v
            t.set(XML_SPACE, "preserve")
            new.append(is_)
        for i, el in enumerate(new):
            c.insert(i, el)

    @staticmethod
    def _sub(parent, tag, text):
        el = etree.SubElement(parent, m(tag))
        el.text = text
        return el

    @staticmethod
    def update_dimension(root):
        dim = root.find(m("dimension"))
        coords = [coordinate_from_string(c.get("r")) for c in root.iter(m("c")) if c.get("r")]
        if dim is None or not coords:
            return
        cols = [column_index_from_string(col) for col, _ in coords]
        rows = [r for _, r in coords]
        dim.set("ref", format_ref(min(cols), min(rows), max(cols), max(rows)))

    def cell_text(self, c):
        if c is None:
            return ""
        if c.get("t") == "s":
            v = c.find(m("v"))
            si = self.sst[int(v.text)] if v is not None and int(v.text) < len(self.sst) else None
            return "".join(t.text or "" for t in si.iter(m("t"))) if si is not None else ""
        if c.get("t") == "inlineStr":
            return "".join(t.text or "" for t in c.iter(m("t")))
        v = c.find(m("v"))
        return v.text or "" if v is not None else ""

    # --- charts, images and comments hanging off a sheet ---

    def merge_sheet_object(self, oname, key, bpart, opart, tpart):
        """Whether a sheet has a drawing / comments at all. What's inside
        them merges separately, object by object."""
        b = self.b.sheet_object(bpart, key)[0] is not None
        o_part = self.o.sheet_object(opart, key)[0]
        t_part, ttype = self.t.sheet_object(tpart, key)
        o, t = o_part is not None, t_part is not None
        if t == b:
            return
        if o != b:  # both branches added one
            self.conflict(oname, f"({key})", "", "added", "added; theirs not merged")
            return
        part = self.sheet_map()[oname][1]
        if not t and self.changed_since_base(o_part):
            self.conflict(oname, f"({key})", "", "changed", "removed; kept ours")
            return
        root = self.tree(part)
        old_el = root.find(m(key)) if key != "comments" else None
        for r in list(self.rels_tree(part)):
            if (old_el is not None and r.get("Id") == old_el.get(RID)) or \
                    (key == "comments" and kind_of(r.get("Type")) == "comments"):
                self.rels_tree(part).remove(r)
        if old_el is not None:
            root.remove(old_el)
        if not t:
            self.objects_taken.append(f"{oname}: {key} removed as on their branch")
            return
        rid = self.add_rel(part, ttype, self.import_part(t_part))
        if key != "comments":
            el = etree.SubElement(root, m(key))
            el.set(RID, rid)
            self.place(root, el, WS_ORDER)
        self.objects_taken.append(f"{oname}: {key} from their branch")

    @staticmethod
    def place(root, el, order):
        """Move el to its schema-mandated position among root's children.
        Extension elements (mc:AlternateContent etc.) are assumed to sit where
        Excel writes them: after workbookPr, or among a sheet's oleObjects."""
        idx = order.index(local(el))
        unknown = order.index("workbookPr" if "workbookPr" in order else "oleObjects")
        for child in root:
            if child is el:
                continue
            name = local(child)
            pos = order.index(name) if name in order else unknown
            if pos > idx:
                child.addprevious(el)
                return

    # --- tables and pivot tables on a sheet ---

    def merge_sheet_collection(self, oname, kind, bpart, opart, tpart):
        """Tables / pivot tables added or removed on their branch."""
        def keys_on(keys, pkg, part):
            if not part:
                return {}
            return {keys.key_of[t]: t for _, typ, t, ext in pkg.rels(part)
                    if not ext and kind_of(typ) == kind and t in keys.key_of}
        b = keys_on(self.kb, self.b, bpart)
        o = keys_on(self.ko, self.o, opart)
        t = keys_on(self.kt, self.t, tpart)
        for key in sorted(set(t) - set(b), key=repr):
            if key in o:
                if normalized(self.ko, o[key]) != normalized(self.kt, t[key]):
                    self.conflict(oname, f"({kind})", "", "added", "added with the same identity; kept ours")
                continue
            self.wire_new(oname, kind, t[key])
        for key in sorted((set(b) - set(t)) & set(o), key=repr):
            if kind == "pivotTable":  # a refresh on our side isn't a reason to keep it
                changed = normalized(self.ko, o[key], True) != normalized(self.kb, b[key], True)
            else:
                changed = self.changed_since_base(o[key])
            if changed:
                self.conflict(oname, f"({kind} {self.o.xml(o[key]).get('name')})", "", "changed", "deleted; kept ours")
            else:
                self.unwire(oname, o[key])
                self.objects_taken.append(f"{oname}: {kind} removed as on their branch")

    def wire_new(self, oname, kind, their_part):
        part = self.sheet_map()[oname][1]
        theirs = self.t.xml(their_part)
        label = theirs.get("displayName") or theirs.get("name")
        new_ref = theirs.get("ref") if kind == "table" else theirs.find(m("location")).get("ref")
        names = set()
        for _, typ, tgt, ext in self.rels_of(part):
            if ext or tgt not in self.res or kind_of(typ) not in ("table", "pivotTable"):
                continue
            x = self.tree(tgt)
            ref = x.get("ref") if kind_of(typ) == "table" else x.find(m("location")).get("ref")
            if overlaps(ref, new_ref):
                self.conflict(oname, f"({kind} {label})", "", "", f"added at {new_ref}, overlapping ours; not merged")
                return
            names.add((x.get("name") or "").lower())
        if kind == "table" and self.t.rels(their_part):
            self.conflict(oname, f"(table {label})", "", "", "added with an external data connection; not merged")
            return
        target = self.import_part(their_part)
        typ = REL_TABLE if kind == "table" else REL_PIVOT
        rid = self.add_rel(part, typ, target)
        if kind == "table":
            root = self.tree(part)
            parts = root.find(m("tableParts"))
            if parts is None:
                parts = etree.SubElement(root, m("tableParts"))
                self.place(root, parts, WS_ORDER)
            etree.SubElement(parts, m("tablePart")).set(RID, rid)
            parts.set("count", str(len(parts)))
        elif label.lower() in names:
            n = 2
            while f"{label} ({n})".lower() in names:
                n += 1
            self.tree(target).set("name", f"{label} ({n})")
        self.objects_taken.append(f"{oname}: {kind} {label} from their branch")

    def unwire(self, oname, our_part):
        part = self.sheet_map()[oname][1]
        root = self.tree(part)
        for rid, typ, tgt, ext in self.rels_of(part):
            if tgt != our_part:
                continue
            self.remove_rel(part, rid)
            for el in [e for e in root.iter() if e.get(RID) == rid]:
                parent = el.getparent()
                parent.remove(el)
                if local(parent) == "tableParts":
                    parent.set("count", str(len(parent)))
                    if len(parent) == 0:
                        root.remove(parent)

    def plan_table_renames(self):
        """Table names are unique per workbook. If both branches added a table
        called Table2, theirs becomes Table3 and their formulas follow."""
        taken = {n.lower() for n, _ in self._names(self.o, {})}
        for key, part in self.ko.part_of.items():
            if key[0] == "table":
                taken.add((self.o.xml(part).get("displayName") or "").lower())
        new = []
        for key, part in self.kt.part_of.items():
            if key[0] == "table":
                name = self.t.xml(part).get("displayName") or ""
                if key in self.kb.part_of:
                    taken.add(name.lower())
                else:
                    new.append(name)
        for name in sorted(new):
            if name.lower() in taken:
                stem, num = re.match(r"(.*?)(\d*)$", name).groups()
                n = int(num or 1) + 1
                while f"{stem}{n}".lower() in taken:
                    n += 1
                self.table_renames[name] = f"{stem}{n}"
                self.notes.append(f"their new table {name} renamed to {stem}{n} (ours has a {name})")
                name = f"{stem}{n}"
            taken.add(name.lower())
        if self.table_renames:
            for cells in self.tc.values():
                for coord, v in cells.items():
                    if isinstance(v, str) and not isinstance(v, Text) and v.startswith("="):
                        cells[coord] = rename_refs(v, self.table_renames)

    # --- whole sheets ---

    def merge_sheet_name(self, bname, oname, tname):
        if tname in (bname, oname):
            return
        if oname != bname:
            self.conflict(oname, "(sheet name)", bname, oname, tname)
        elif tname in self.sheet_map():
            self.conflict(oname, "(sheet name)", bname, oname, f"{tname} (name already taken)")
        else:
            self.sheet_el[oname].set("name", tname)

    def delete_sheet(self, oname):
        el = self.sheet_el[oname]
        self.remove_rel(self.wb, el.get(RID))
        el.getparent().remove(el)

    def add_sheet(self, name, part):
        sheets = self.tree(self.wb).find(m("sheets"))
        sid = max([int(s.get("sheetId")) for s in sheets] + [0]) + 1
        el = etree.SubElement(sheets, m("sheet"))
        el.set("name", name)
        el.set("sheetId", str(sid))
        el.set(RID, self.add_rel(self.wb, WORKSHEET_REL, part))
        return el

    def import_sheet(self, tname, tpart):
        """A sheet only they added: copy it whole, charts, tables and pivots
        included. Shared strings become inline strings since the string table
        isn't merged."""
        root = self.t.xml(tpart)
        sst = self.t.shared_strings()
        for c in root.iter(m("c")):
            for a in ("cm", "vm"):
                c.attrib.pop(a, None)
            if c.get("t") == "s":
                v = c.find(m("v"))
                si = sst[int(v.text)]
                c.remove(v)
                c.set("t", "inlineStr")
                is_ = etree.SubElement(c, m("is"))
                for ch in si:
                    is_.append(copy.deepcopy(ch))
        if self.table_renames:
            for f in root.iter(m("f")):
                if f.text:
                    f.text = rename_refs(f.text, self.table_renames)
        if not self.same_styles:  # their style numbers mean nothing in our styles.xml
            for el in root.iter(m("c"), m("row"), m("col")):
                for a in ("s", "style", "customFormat"):
                    el.attrib.pop(a, None)
            for tag in ("conditionalFormatting", "extLst"):
                for el in root.findall(m(tag)):
                    root.remove(el)
        ours = tpart if tpart not in self.res else self.unique_name(tpart)
        self.imported[tpart] = ours
        self.put(ours, b"")
        self.trees[ours] = root
        self.copy_content_type(tpart, ours)
        keep = ("drawing", "vmlDrawing", "comments", "hyperlink", "printerSettings", "table", "pivotTable")
        dropped = self.copy_rels(tpart, ours, keep_rel=lambda typ: kind_of(typ) in keep)
        for rid, typ in dropped:
            for el in [e for e in root.iter() if e.get(RID) == rid]:
                el.getparent().remove(el)
            self.conflict(tname, f"({kind_of(typ)})", "", "", "on their new sheet; not merged")
        self.new_sheet_el[tname] = self.add_sheet(tname, ours)
        self.objects_taken.append(f"{tname}: new sheet from their branch")

    # --- defined names ---

    def _names(self, pkg, lineage):
        wb = pkg.xml(pkg.workbook)
        if wb is None:
            return {}
        sheets = [s.get("name") for s in wb.find(m("sheets"))]
        dn = wb.find(m("definedNames"))
        out = {}
        for d in (dn if dn is not None else []):
            lid = d.get("localSheetId")
            scope = sheets[int(lid)] if lid is not None and int(lid) < len(sheets) else None
            scope = lineage.get(scope, scope) if scope else None
            attrs = tuple(sorted((k, v) for k, v in d.attrib.items() if k != "localSheetId"))
            out[(d.get("name"), scope)] = (d.text, attrs)
        return out

    def capture_names(self):
        wb = self.tree(self.wb)
        sheets = list(wb.find(m("sheets")))
        self.name_els = {}
        dn = wb.find(m("definedNames"))
        for d in (dn if dn is not None else []):
            lid = d.get("localSheetId")
            scope_el = sheets[int(lid)] if lid is not None and int(lid) < len(sheets) else None
            scope = self.lineage_o.get(scope_el.get("name"), scope_el.get("name")) if scope_el is not None else None
            self.name_els[(d.get("name"), scope)] = (d, scope_el)

    def scope_element(self, scope):
        if scope is None:
            return None
        ours = {v: k for k, v in self.lineage_o.items()}.get(scope, scope)
        return self.sheet_el.get(ours) or self.new_sheet_el.get(scope)

    def merge_names(self):
        b = self._names(self.b, {})
        o = self._names(self.o, self.lineage_o)
        t = self._names(self.t, self.lineage_t)
        wb = self.tree(self.wb)
        moved = lambda v, maps: v and (rewrite_refs("=" + v[0], None, maps)[1:] if v[0] else v[0], v[1])
        for key in set(b) | set(t):
            bv, ov, tv = moved(b.get(key), self.maps_b), o.get(key), moved(t.get(key), self.maps_t)
            if tv == bv or tv == ov:
                continue
            if ov != bv:
                self.conflict("(names)", key[0], bv and bv[0], ov and ov[0], tv and tv[0])
                continue
            if tv is None:
                d, _ = self.name_els.pop(key)
                d.getparent().remove(d)
                continue
            if key in self.name_els:
                d = self.name_els[key][0]
            elif key[1] is not None and self.scope_element(key[1]) is None:
                self.conflict("(names)", key[0], None, None, f"{tv[0]} (its sheet wasn't merged)")
                continue
            else:
                dn = wb.find(m("definedNames"))
                if dn is None:
                    dn = etree.SubElement(wb, m("definedNames"))
                    self.place(wb, dn, WB_ORDER)
                d = etree.SubElement(dn, m("definedName"))
                self.name_els[key] = (d, self.scope_element(key[1]))
            for k in list(d.attrib):
                if k != "localSheetId":
                    del d.attrib[k]
            for k, v in tv[1]:
                d.set(k, v)
            d.text = rename_refs(tv[0], self.table_renames) if tv[0] else tv[0]

    # --- charts, drawings, images, comments, macros: object by object ---

    def family(self, key):
        while key and key[0] in ("sub", "shape"):
            key = key[1]
        return key[0] if key else ""

    def merge_object_parts(self):
        for key, bpart in sorted(self.kb.part_of.items(), key=lambda kv: repr(kv[0])):
            if key[0] in ("pkg", "workbook", "sheet", "part") or (key[0] == "wb" and key[1] in WB_KEEP):
                continue
            if self.family(key) in ("table", "pivot", "cache"):
                continue  # merged by merge_tables / merge_pivots
            opart, tpart = self.ko.part_of.get(key), self.kt.part_of.get(key)
            b = normalized(self.kb, bpart)
            o, t = normalized(self.ko, opart), normalized(self.kt, tpart)
            if t is None or t == b or t == o:
                continue  # a part they deleted vanishes once nothing links to it
            if o is None:
                self.conflict("(file)", tpart, "", "", "deleted on ours, changed on theirs; kept deletion")
            elif o != b:
                self.conflict("(file)", opart, "", "", "changed on both branches; kept ours")
            else:
                self.take_theirs(opart, tpart)
                self.objects_taken.append(opart)

    # --- tables ---

    def merge_tables(self):
        for key, bpart in sorted(self.kb.part_of.items(), key=lambda kv: repr(kv[0])):
            if key[0] != "table":
                continue
            opart, tpart = self.ko.part_of.get(key), self.kt.part_of.get(key)
            if opart is None or tpart is None:
                if opart is None and tpart and normalized(self.kt, tpart) != normalized(self.kb, bpart):
                    self.conflict("(table)", self.t.xml(tpart).get("displayName"), "", "deleted",
                                  "changed; kept deletion")
                continue  # removal on their side is handled per sheet
            b, o, t = (normalized(k, p) for k, p in ((self.kb, bpart), (self.ko, opart), (self.kt, tpart)))
            if t == b or t == o:
                continue
            name = self.o.xml(opart).get("displayName")
            if o == b:
                self.put(opart, self.t.parts[tpart])
                if not self.same_styles:
                    strip_dxf(self.tree(opart))
            else:
                merged = self.merge_table_xml(self.b.xml(bpart), self.o.xml(opart), self.t.xml(tpart))
                if merged is None:
                    self.conflict("(table)", name, "", "", "changed on both branches; kept ours")
                    continue
                self.trees[opart] = merged
                self.put(opart, b"")
            self.touched_tables.add(opart)
            self.objects_taken.append(f"table {name}")

    def merge_table_xml(self, b, o, t):
        """Field-by-field merge of a table definition: its range, each column,
        filter and style. None if the two sides can't be reconciled."""
        r = copy.deepcopy(o)
        their_el = lambda el: strip_dxf(el) or el if not self.same_styles else el
        for k in set(b.attrib) | set(t.attrib):
            if k in ("id", "ref") or (not self.same_styles and k.lower().endswith("dxfid")):
                continue
            v, ok = _three_way(b.get(k), o.get(k), t.get(k))
            if not ok:
                return None
            r.set(k, v) if v is not None else r.attrib.pop(k, None)
        ref, ok = _three_way(b.get("ref"), o.get("ref"), t.get("ref"))
        if not ok:  # e.g. we added rows, they added a column: merge each edge
            edges = [_three_way(*x) for x in zip(*(parse_ref(x.get("ref")) for x in (b, o, t)))]
            if not all(ok for _, ok in edges):
                return None
            ref = format_ref(*(v for v, _ in edges))
        r.set("ref", ref)

        cols = lambda x: x.find(m("tableColumns"))
        ids = lambda x: [c.get("id") for c in cols(x)]
        order, ok = _three_way(ids(b), ids(o), ids(t))
        if not ok:
            return None
        by_id = lambda x: {c.get("id"): c for c in cols(x)}
        bc, oc, tc = by_id(b), by_id(o), by_id(t)
        ser = lambda el: etree.tostring(el) if el is not None else None
        new_cols = []
        for cid in order:
            if cid not in oc:
                new_cols.append(their_el(copy.deepcopy(tc[cid])))
                continue
            if cid not in tc or cid not in bc:
                new_cols.append(copy.deepcopy(oc[cid]))
                continue
            _, ok = _three_way(ser(bc[cid]), ser(oc[cid]), ser(tc[cid]))
            if not ok:
                return None
            pick = oc[cid] if ser(tc[cid]) in (ser(bc[cid]), ser(oc[cid])) else their_el(copy.deepcopy(tc[cid]))
            new_cols.append(copy.deepcopy(pick))
        rc = cols(r)
        for c in list(rc):
            rc.remove(c)
        for c in new_cols:
            rc.append(c)
        rc.set("count", str(len(new_cols)))

        for tag in ("autoFilter", "sortState", "tableStyleInfo", "extLst"):
            def plain(x):
                el = x.find(m(tag))
                if el is None:
                    return None
                el = copy.deepcopy(el)
                el.attrib.pop("ref", None)
                return etree.tostring(el)
            _, ok = _three_way(plain(b), plain(o), plain(t))
            if not ok:
                return None
            if plain(t) not in (plain(b), plain(o)):
                old = r.find(m(tag))
                if old is not None:
                    r.remove(old)
                if t.find(m(tag)) is not None:
                    el = their_el(copy.deepcopy(t.find(m(tag))))
                    r.append(el)
                    self.place(r, el, TABLE_ORDER)
        c1, r1, c2, r2 = parse_ref(r.get("ref"))
        if c2 - c1 + 1 != len(new_cols):
            return None
        af = r.find(m("autoFilter"))
        if af is not None:
            totals = int(r.get("totalsRowCount") or 0)
            af.set("ref", format_ref(c1, r1, c2, r2 - totals))
        ss = r.find(m("sortState"))
        if ss is not None and ss.get("ref") and not overlaps(ss.get("ref"), r.get("ref")):
            r.remove(ss)
        return r

    def fix_tables(self):
        """Every table needs a workbook-unique id and name, and its column
        names must match its header cells or Excel 'repairs' the file."""
        live = []
        for name, (_, spart) in self.sheet_map().items():
            for _, typ, tgt, ext in self.rels_of(spart):
                if not ext and kind_of(typ) == "table" and tgt in self.res:
                    live.append((tgt, spart))
        imported = set(self.imported.values())
        live.sort(key=lambda x: x[0] in imported)  # ours keep their ids
        used = set()
        top = max([int(self.tree(t).get("id") or 0) for t, _ in live] + [0])
        for tpart, spart in live:
            root = self.tree(tpart)
            if root.get("id") in used:
                top += 1
                root.set("id", str(top))
            used.add(root.get("id"))
            if tpart in imported:
                self.touched_tables.add(tpart)
                old = root.get("displayName")
                if old in self.table_renames:
                    root.set("displayName", self.table_renames[old])
                    root.set("name", self.table_renames[old])
                for f in root.iter(m("calculatedColumnFormula")):
                    if f.text:
                        f.text = rename_refs(f.text, self.table_renames)
                if not self.same_styles:
                    strip_dxf(root)
            if (tpart in self.touched_tables or spart in self.trees) and root.get("headerRowCount") != "0":
                c1, r1, _, _ = parse_ref(root.get("ref"))
                cells = dict(iter_sheet_cells(self.tree(spart)))
                cols = list(root.find(m("tableColumns")))
                names = [self.cell_text(cells.get(f"{get_column_letter(c1 + i)}{r1}")) or xdecode(col.get("name"))
                         for i, col in enumerate(cols)]
                if len({n.lower() for n in names}) == len(names):
                    for col, n in zip(cols, names):
                        if xdecode(col.get("name")) != n:
                            col.set("name", xencode(n))

    # --- pivot tables ---

    def merge_pivots(self):
        """A pivot table, its cache and the cache's records only make sense
        together, so they merge as one bundle per cache. A side that only
        refreshed (same layout, new data) doesn't count as an edit: Excel is
        told to refresh on open, which rebuilds the pivot from merged data."""
        for ckey, bcache in sorted(self.kb.part_of.items(), key=lambda kv: repr(kv[0])):
            if ckey[0] != "cache":
                continue
            members = [ckey] + [k for k in self.kb.part_of if k[0] == "sub" and k[1] == ckey]
            members += [pk for pk, pp in self.kb.pivots if self.b.target(pp, "pivotCacheDefinition") == bcache]
            bundle = lambda keys, derived: tuple(normalized(keys, keys.part_of.get(k), derived) for k in members
                                                 if not (derived and k[0] == "sub"))  # records are pure data
            b, o, t = bundle(self.kb, False), bundle(self.ko, False), bundle(self.kt, False)
            if t == b or t == o:
                continue
            names = ", ".join(str(k[2]) for k in members if k[0] == "pivot")
            if o != b:
                if bundle(self.kt, True) == bundle(self.kb, True):
                    self.stale_pivots = True  # they only refreshed
                    continue
                if bundle(self.ko, True) != bundle(self.kb, True):
                    self.conflict("(pivot)", names, "", "", "changed on both branches; kept ours")
                    self.stale_pivots = True
                    continue
            for k in members:
                opart, tpart = self.ko.part_of.get(k), self.kt.part_of.get(k)
                if opart and tpart:
                    self.take_theirs(opart, tpart)
            self.stale_pivots = True
            self.objects_taken.append(f"pivot table(s) {names}")

    def fix_pivots(self):
        """Register every pivot's cache in workbook.xml with a matching cacheId,
        drop caches nothing uses any more, and have Excel refresh pivots whose
        source data changed in the merge."""
        wb = self.tree(self.wb)
        pcs = wb.find(m("pivotCaches"))
        wb_caches = {r.get("Id"): resolve(self.wb, r.get("Target")) for r in self.rels_tree(self.wb)
                     if kind_of(r.get("Type")) == "pivotCacheDefinition"}
        entry = {wb_caches.get(pc.get(RID)): pc for pc in (pcs if pcs is not None else [])}
        used = set()
        imported = set(self.imported.values())
        for name, (_, spart) in self.sheet_map().items():
            for _, typ, pv, ext in self.rels_of(spart):
                if ext or kind_of(typ) != "pivotTable" or pv not in self.res:
                    continue
                cache = next((t for _, ty, t, e in self.rels_of(pv) if kind_of(ty) == "pivotCacheDefinition"), None)
                if cache not in self.res:
                    continue
                if cache not in entry:
                    if pcs is None:
                        pcs = etree.SubElement(wb, m("pivotCaches"))
                        self.place(wb, pcs, WB_ORDER)
                    pc = etree.SubElement(pcs, m("pivotCache"))
                    pc.set("cacheId", str(max([int(p.get("cacheId")) for p in pcs if p.get("cacheId")] + [0]) + 1))
                    pc.set(RID, self.add_rel(self.wb, REL_CACHE, cache))
                    entry[cache] = pc
                root = self.tree(pv)
                if root.get("cacheId") != entry[cache].get("cacheId"):
                    root.set("cacheId", entry[cache].get("cacheId"))
                if pv in imported and not self.same_styles:
                    strip_dxf(root)
                if cache in imported and self.table_renames:
                    for src in self.tree(cache).iter(m("worksheetSource")):
                        if src.get("name") in self.table_renames:
                            src.set("name", self.table_renames[src.get("name")])
                used.add(cache)
        ours_used = {self.o.target(p, "pivotCacheDefinition") for _, p in self.ko.pivots}
        for cache, pc in list(entry.items()):
            if cache not in used and cache in ours_used:
                self.remove_rel(self.wb, pc.get(RID))
                pcs.remove(pc)
        if pcs is not None and len(pcs) == 0:
            wb.remove(pcs)
        if self.cells_taken or self.stale_pivots:
            for cache in used:
                root = self.tree(cache)
                src = root.find(m("cacheSource"))
                if src is not None and src.get("type") == "worksheet":
                    root.set("refreshOnLoad", "1")

    def report_unmerged_additions(self):
        for part in sorted(reachable(self.t.parts)):
            key = self.kt.key_of.get(part)
            if key in self.kb.part_of or part in self.imported or NOT_OBJECT_PART.match(part):
                continue
            self.conflict("(file)", part, "", "", "added on their branch; not merged")

    def write_conflict_sheet(self):
        if not self.conflicts:
            return
        root = etree.Element(m("worksheet"), nsmap={None: MAIN, "r": REL})
        sd = etree.SubElement(root, m("sheetData"))
        header = ["sheet", "cell", "base", "ours (kept)", "theirs", "go to"]
        live = self.sheet_map()
        for i, row in enumerate([header] + [list(c) for c in self.conflicts], start=1):
            r = etree.SubElement(sd, m("row"))
            r.set("r", str(i))
            is_cell = i > 1 and re.match(r"^[A-Z]+\d+$", str(row[1]))
            values = [fmt(v) if is_cell and j in (2, 3, 4) else ("" if v is None else str(v))
                      for j, v in enumerate(row)]
            if is_cell and row[0] in live:
                sheet = row[0].replace("'", "''")
                values.append(f"=HYPERLINK(\"#'{sheet}'!{row[1]}\",\"{row[0]}!{row[1]}\")")
            for j, v in enumerate(values):
                c = etree.SubElement(r, m("c"))
                c.set("r", f"{get_column_letter(j + 1)}{i}")
                self.write_value(c, v, as_text=j < 5)
        if CONFLICT_SHEET in live:
            self.trees[live[CONFLICT_SHEET][1]] = root
            return
        part = self.unique_name("xl/worksheets/sheet1.xml")
        self.put(part, b"")
        self.trees[part] = root
        ct = self.tree("[Content_Types].xml")
        o = etree.SubElement(ct, f"{{{CTYPES}}}Override")
        o.set("PartName", "/" + part)
        o.set("ContentType", WORKSHEET_CT)
        self.add_sheet(CONFLICT_SHEET, part)

    # --- write out ---

    def finish(self):
        wb = self.tree(self.wb)
        sheets = list(wb.find(m("sheets")))
        for key, (d, scope_el) in list(self.name_els.items()):
            if scope_el is None:
                d.attrib.pop("localSheetId", None)
            elif scope_el.getparent() is None:  # its sheet was deleted
                d.getparent().remove(d)
            else:
                d.set("localSheetId", str(sheets.index(scope_el)))
        dn = wb.find(m("definedNames"))
        if dn is not None and len(dn) == 0:
            wb.remove(dn)
        for view in wb.iter(m("workbookView")):
            for a in ("activeTab", "firstSheet"):
                if view.get(a) and int(view.get(a)) >= len(sheets):
                    view.set(a, "0")
        # Cached formula results are stale now; have Excel recalculate on open
        # and rebuild its calculation chain.
        calc = wb.find(m("calcPr"))
        if calc is None:
            calc = etree.SubElement(wb, m("calcPr"))
            self.place(wb, calc, WB_ORDER)
        calc.set("fullCalcOnLoad", "1")
        for r in list(self.rels_tree(self.wb)):
            if kind_of(r.get("Type")) == "calcChain":
                self.rels_tree(self.wb).remove(r)

        for name, root in self.trees.items():
            self.res[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        # Drop parts nothing links to any more (deleted sheets, replaced charts).
        before = reachable(self.o.parts) | set(self.imported.values())
        after = reachable(self.res)
        for name in list(self.res):
            if name != "[Content_Types].xml" and not name.endswith(".rels") and name in before and name not in after:
                self.drop(name)
                self.drop(rels_name(name))
        ct = self.tree("[Content_Types].xml")
        self.res["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)

        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            names = ["[Content_Types].xml"] + [n for n in self.order if n != "[Content_Types].xml"]
            for name in names:
                if name in self.res:
                    z.writestr(name, self.res[name])
        return out.getvalue()


def merge(base_path, ours_path, theirs_path, display_path=None):
    """Write the merge of THEIRS into OURS over OURS (git's convention).
    Exit 0 = clean, 1 = conflicts (listed in a _merge_conflicts sheet)."""
    name = display_path or ours_path
    try:
        mg = Merger(base_path, ours_path, theirs_path)
        data = mg.run()
    except Exception as e:  # never leave a half-written workbook behind
        say(f"xlgit couldn't merge {name} automatically ({_reason(e)}).\n"
            f"  Your version is kept, unchanged. To take theirs instead:\n"
            f"    git checkout --theirs -- \"{name}\"\n"
            f"  Then: git add \"{name}\"  and  git commit\n"
            f"  Please report this at {ISSUES} so it can be fixed. To share the files safely:\n"
            f"    xlgit scrub --merge \"{name}\"")
        return 1
    with open(ours_path, "wb") as f:
        f.write(data)
    took = f"{mg.cells_taken} cell(s)" + (f" and {len(mg.objects_taken)} object(s)" if mg.objects_taken else "")
    if mg.swapped:  # built on their version; the cells written were yours
        took = took.replace("cell(s)", "of your edited cell(s)", 1) + " onto their inserted/deleted rows"
    if not mg.conflicts:
        say(f"xlgit merged {name}: " + (f"moved {took}" if mg.swapped else f"took {took} from the other branch")
            + ", no conflicts.")
    for note in mg.notes:
        say(f"  note: {note}")
    if mg.conflicts:
        say(f"xlgit merged {name}: " + (f"moved {took}" if mg.swapped else f"took {took} from the other branch")
            + ", but "
            f"{len(mg.conflicts)} change(s) clash.")
        for sheet, coord, bv, ov, tv in mg.conflicts[:20]:
            say(f"  {sheet} {coord}: yours {fmt(ov)}, theirs {fmt(tv)} (was {fmt(bv)})")
        if len(mg.conflicts) > 20:
            say(f"  ...and {len(mg.conflicts) - 20} more.")
        say(f"  Your values were kept. Every clash is listed, with a link, on the sheet {CONFLICT_SHEET!r}.\n"
            f"  To finish: open {name}, fix those cells, delete the {CONFLICT_SHEET} sheet, save, then\n"
            f"    git add \"{name}\"\n"
            f"    git commit")
    return 1 if mg.conflicts else 0


def say(msg):
    print(msg, file=sys.stderr)


def _reason(e):
    if isinstance(e, zipfile.BadZipFile):
        return "it isn't an .xlsx/.xlsm workbook; old .xls files aren't supported, save as .xlsx"
    if isinstance(e, etree.XMLSyntaxError):
        return f"the workbook is damaged: {e}"
    return f"{type(e).__name__}: {e}"


# ---------- scrub: shareable copies for bug reports ----------

_TINY = {  # 1x1 images that stand in for the originals
    "png": bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                         "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"),
    "gif": bytes.fromhex("47494638396101000100800000ffffff00000021f90401000000002c00000000010001000002024401003b"),
    "jpeg": bytes.fromhex("ffd8ffe000104a46494600010100000100010000ffdb004300080606070605080707070909080a0c140d0c0b0b0c"
                          "1912130f141d1a1f1e1d1a1c1c20242e2720222c231c1c2837292c30313434341f27393d38323c2e333432ffc000"
                          "0b080001000101011100ffc4001f0000010501010101010100000000000000000102030405060708090a0bffc400"
                          "b5100002010303020403050504040000017d01020300041105122131410613516107227114328191a1082342b1c1"
                          "1552d1f02433627282090a161718191a25262728292a3435363738393a434445464748494a535455565758595a"
                          "636465666768696a737475767778797a838485868788898a92939495969798999aa2a3a4a5a6a7a8a9aab2b3b4b5"
                          "b6b7b8b9bac2c3c4c5c6c7c8c9cad2d3d4d5d6d7d8d9dae1e2e3e4e5e6e7e8e9eaf1f2f3f4f5f6f7f8f9faffda00"
                          "080101000003f00fbfffd9"),
}
_NUM = re.compile(r"^(-?)(\d+)(?:\.(\d+))?([eE][-+]?\d+)?$")
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


class Scrubber:
    """Replaces every piece of text and every number in workbooks with made-up
    ones, keeping their structure, so a workbook that trips xlgit up can be
    shared in a bug report. Equal values stay equal (across all the files
    scrubbed together), so a merge of scrubbed versions behaves the same."""

    def __init__(self, key=None):
        self.key = key or os.urandom(16)
        self.kept = set()

    def _rng(self, kind, value):
        import hmac
        import random
        return random.Random(hmac.new(self.key, f"{kind}\0{value}".encode(), "sha256").digest())

    def text(self, s):
        if not s or s in ERRORS or s.upper() in ("TRUE", "FALSE"):
            return s
        rng = self._rng("t", s)
        lower, upper = "abcdefghijklmnopqrstuvwxyz", "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
        return "".join(rng.choice(lower) if ch.islower() else rng.choice(upper) if ch.isupper()
                       else rng.choice("0123456789") if ch.isdigit()
                       else rng.choice(lower) if ch.isalpha() else ch for ch in s)

    def number(self, text, date=False):
        try:
            value = float(text)
        except (TypeError, ValueError):
            return text
        if value == 0 or value != value:
            return text
        # Numbers xlgit treats as equal (to 15 digits) get the same stand-in.
        text = f"{value:.15g}"
        mt = _NUM.match(text)
        if not mt:
            return text
        sign, whole, frac, exp = mt.groups()
        rng = self._rng("n", text)
        if date:  # somewhere in 1998-2028, keeping any time of day
            whole = str(rng.randint(36000, 47000))
        else:
            whole = "".join(rng.choice("123456789" if i == 0 and len(whole) > 1 else "0123456789")
                            for i in range(len(whole)))
        frac = "".join(rng.choice("0123456789") for _ in frac or "")
        if frac and not frac.strip("0"):
            frac = frac[:-1] + "5"
        return f"{sign}{whole}{'.' + frac if frac else ''}{exp or ''}"

    def formula(self, f):
        """Text inside quotes in a formula is data too."""
        if '"' not in f:
            return f
        try:
            from openpyxl.formula.tokenizer import Tokenizer
            tok = Tokenizer("=" + f)
        except Exception:
            return re.sub(r'"((?:[^"]|"")*)"', lambda mt: '"' + self.text(mt.group(1).replace('""', '"'))
                          .replace('"', '""') + '"', f)
        for t in tok.items:
            if t.type == "OPERAND" and t.subtype == "TEXT":
                inner = t.value[1:-1].replace('""', '"')
                t.value = '"' + self.text(inner).replace('"', '""') + '"'
        return tok.render()[1:]

    def _texts(self, root, tags):
        for el in root.iter(*tags):
            if el.text:
                el.text = self.text(el.text)

    def scrub(self, data):
        pkg = Package(data)
        dates = _date_styles(pkg)
        out, dropped = {}, set()
        for name in pkg.order:
            part = pkg.parts[name]
            lower = name.lower()
            ext = lower.rsplit(".", 1)[-1]
            if lower.endswith("vbaproject.bin") or lower.endswith("vbadata.xml"):
                dropped.add(name)
                self.kept.add("macros removed")
                continue
            if lower.startswith("xl/media/") or "/media/" in lower:
                if ext in ("png", "gif") or ext in ("jpg", "jpeg"):
                    out[name] = _TINY["jpeg" if ext in ("jpg", "jpeg") else ext]
                else:
                    out[name] = part
                    self.kept.add(f"images in .{ext} format kept as they are")
                continue
            if not lower.endswith((".xml", ".rels", ".vml")):
                out[name] = part
                continue
            try:
                root = parse_xml(part)
            except etree.XMLSyntaxError:
                out[name] = part
                continue
            self._scrub_part(name, root, pkg, dates)
            out[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        if dropped:
            self._unlink(out, dropped)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name in pkg.order:
                if name in out:
                    z.writestr(name, out[name])
        return buf.getvalue()

    def _unlink(self, out, dropped):
        for name in list(out):
            if name.endswith(".rels"):
                root = parse_xml(out[name])
                owner = name.replace("_rels/", "")[:-5]
                owner = "" if owner == "." else owner
                for r in list(root):
                    if r.get("TargetMode") != "External" and resolve(owner, r.get("Target")) in dropped:
                        root.remove(r)
                out[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = parse_xml(out["[Content_Types].xml"])
        for o in list(ct):
            if (o.get("PartName") or "").lstrip("/") in dropped:
                ct.remove(o)
        out["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _scrub_part(self, name, root, pkg, dates):
        tag = local(root)
        if name.endswith(".rels"):
            for r in root:
                if r.get("TargetMode") == "External":
                    r.set("Target", "https://example.com/" if "hyperlink" in (r.get("Type") or "")
                          else "external.xlsx")
            return
        if tag in ("worksheet", "macrosheet", "externalLink", "dialogsheet"):
            self._scrub_cells(root, dates)
            self._texts(root, (m("oddHeader"), m("oddFooter"), m("evenHeader"), m("evenFooter"),
                               m("firstHeader"), m("firstFooter")))
            for el in root.iter(m("hyperlink")):
                for a in ("display", "tooltip"):
                    if el.get(a):
                        el.set(a, self.text(el.get(a)))
            return
        if tag == "sst":
            for si in root.iter(m("si")):
                for rph in list(si.iter(m("rPh"))):
                    rph.getparent().remove(rph)
                self._texts(si, (m("t"),))
            return
        if tag == "table":
            for col in root.iter(m("tableColumn")):
                col.set("name", xencode(self.text(xdecode(col.get("name")))))
                for a in ("totalsRowLabel",):
                    if col.get(a):
                        col.set(a, self.text(col.get(a)))
            for f in root.iter(m("calculatedColumnFormula"), m("totalsRowFormula")):
                if f.text:
                    f.text = self.formula(f.text)
            return
        if tag == "pivotCacheDefinition":
            for cf in root.iter(m("cacheField")):
                cf.set("name", self.text(cf.get("name")))
            self._scrub_items(root)
            return
        if tag == "pivotCacheRecords":
            self._scrub_items(root)
            return
        if tag == "pivotTableDefinition":
            for el in root.iter(m("item"), m("pivotField"), m("dataField")):
                if el.get("n"):
                    el.set("n", self.text(el.get("n")))
                if local(el) == "dataField" and el.get("name"):
                    el.set("name", self.text(el.get("name")))
            return
        if tag in ("comments", "ThreadedComments", "personList", "chartSpace", "wsDr", "userShapes"):
            for el in root.iter():
                if not isinstance(el.tag, str):
                    continue
                ln = local(el)
                if ln in ("t", "text") and el.text:
                    el.text = self.text(el.text)
                elif ln == "v" and el.text and el.getparent() is not None and local(el.getparent()) == "pt":
                    cache = el.getparent().getparent()
                    el.text = self.number(el.text) if cache is not None and local(cache) == "numCache" \
                        else self.text(el.text)
                if el.get("displayName") and ln == "person":
                    el.set("displayName", self.text(el.get("displayName")))
                if el.get("authorId") is None and ln == "author" and el.text:
                    el.text = self.text(el.text)
            return
        if tag in ("coreProperties", "Properties"):
            for el in root.iter():
                if isinstance(el.tag, str) and local(el) in (
                        "creator", "lastModifiedBy", "title", "subject", "description", "keywords",
                        "category", "Company", "Manager", "HyperlinkBase", "Template"):
                    el.text = ""
            return
        if "vml" in name.lower():
            for el in root.iter():
                if isinstance(el.tag, str) and el.text and el.text.strip() and len(el) == 0 \
                        and local(el) not in ("Anchor", "Row", "Column", "ClientData"):
                    el.text = self.text(el.text)
            return

    def _scrub_items(self, root):
        for el in root.iter(m("s"), m("n"), m("d"), m("e")):
            v = el.get("v")
            if v is None:
                continue
            if local(el) == "s":
                el.set("v", self.text(v))
            elif local(el) == "n":
                el.set("v", self.number(v))
        for el in root.iter(m("sharedItems")):
            for a in ("minValue", "maxValue"):
                if el.get(a):
                    del el.attrib[a]

    def _scrub_cells(self, root, dates):
        for c in root.iter(m("c")):
            t = c.get("t")
            f = c.find(m("f"))
            v = c.find(m("v"))
            if f is not None:
                if f.text:
                    f.text = self.formula(f.text)
                if v is not None:  # a cached result: Excel recalculates it on open
                    c.remove(v)
                if t in ("str", "s", "e", "b"):
                    c.attrib.pop("t", None)
                continue
            if t == "inlineStr":
                is_ = c.find(m("is"))
                if is_ is not None:
                    self._texts(is_, (m("t"),))
            elif t == "str" and v is not None and v.text:
                v.text = self.text(v.text)
            elif t in (None, "n") and v is not None and v.text:
                v.text = self.number(v.text, date=c.get("s") in dates)


def scrub_files(paths, merge_of=None):
    """Write NAME.scrubbed.xlsx next to each file, all with one mapping."""
    sc = Scrubber()
    written = []
    for p in paths:
        with open(p, "rb") as f:
            data = f.read()
        stem, ext = os.path.splitext(p)
        out = f"{stem}.scrubbed{ext or '.xlsx'}"
        with open(out, "wb") as f:
            f.write(sc.scrub(data))
        written.append(out)
    return written, sc


def scrub_merge(path):
    """During a merge with a conflict on PATH: scrub the three versions git
    keeps (base, yours, theirs) into one zip for a bug report."""
    root = repo_root()
    if not root:
        raise UserError("run this inside the repository where the merge happened.")
    rel = os.path.relpath(os.path.abspath(path), root).replace("\\", "/")
    blobs = {}
    for stage, label in ((1, "base"), (2, "ours"), (3, "theirs")):
        r = subprocess.run(["git", "show", f":{stage}:{rel}"], cwd=root, capture_output=True)
        if r.returncode:
            raise UserError(f"git has no {label} version of {rel}; is a merge of it in progress? "
                            "(To scrub files directly: xlgit scrub FILE...)")
        blobs[label] = r.stdout
    sc = Scrubber()
    stem = os.path.splitext(os.path.basename(rel))[0]
    out = os.path.abspath(f"{stem}-merge-report.zip")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for label, data in blobs.items():
            z.writestr(f"{label}.xlsx", sc.scrub(data))
        z.writestr("README.txt", f"xlgit {__version__} merge report for {stem}: base, ours and theirs, scrubbed.\n")
    return out, sc


def _scrub_notice(sc):
    kept = "; ".join(sorted(sc.kept))
    return ("  Every cell value, text box, comment, chart label and file property was replaced with made-up\n"
            "  values (equal values stay equal). Kept as they were: sheet names, named ranges, formulas\n"
            "  (text in quotes replaced), formatting and layout" + (f"; {kept}" if kept else "") + ".\n"
            "  Open it and check before you share it.")


# ---------- setup ----------

ISSUES = "https://github.com/MichaelFowler1/excel-git/issues"
ACTION = "MichaelFowler1/excel-git"
ATTR_LINES = [f"{ext} diff=xlsx merge=xlsx" for ext in EXTS]
WORKFLOW = f"""\
# Posts a cell-by-cell diff of changed Excel workbooks on every pull request.
# Made by `xlgit install --github`; see https://github.com/{ACTION}
name: Excel diff

on:
  pull_request:
    paths: ["**/*.xlsx", "**/*.xlsm", "**/*.XLSX", "**/*.XLSM"]

permissions:
  contents: read
  pull-requests: write

jobs:
  excel-diff:
    runs-on: ubuntu-latest
    steps:
      - uses: actions/checkout@v4
        with:
          fetch-depth: 0
      - uses: {ACTION}@v{__version__}
"""


def git(*args, check=True):
    r = subprocess.run(["git", *args], capture_output=True)
    if check and r.returncode:
        raise UserError(r.stderr.decode(errors="replace").strip() or f"git {' '.join(args)} failed")
    return r.stdout


class UserError(Exception):
    """A problem the user can fix; shown without a traceback."""


def repo_root():
    r = subprocess.run(["git", "rev-parse", "--show-toplevel"], capture_output=True, text=True)
    return r.stdout.strip() if r.returncode == 0 else None


def global_attributes_file():
    """The attributes file git reads for every repository."""
    configured = git("config", "--global", "--get", "core.attributesFile", check=False).decode().strip()
    if configured:
        return os.path.expanduser(configured)
    xdg = os.environ.get("XDG_CONFIG_HOME") or os.path.join(os.path.expanduser("~"), ".config")
    return os.path.join(xdg, "git", "attributes")


def _add_lines(path, lines):
    existing = open(path, encoding="utf-8").read() if os.path.exists(path) else ""
    missing = [ln for ln in lines if ln not in existing.splitlines()]
    if missing:
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        with open(path, "a", encoding="utf-8") as f:
            if existing and not existing.endswith("\n"):
                f.write("\n")
            f.write("\n".join(missing) + "\n")


def _remove_lines(path, lines):
    if os.path.exists(path):
        kept = [ln for ln in open(path, encoding="utf-8").read().splitlines() if ln not in lines]
        with open(path, "w", encoding="utf-8") as f:
            f.write("".join(ln + "\n" for ln in kept))


def _set_drivers(where):
    """Point git's diff and merge drivers for workbooks at this program."""
    if getattr(sys, "frozen", False):  # the standalone download: git runs the program itself
        cmd = '"' + sys.executable.replace("\\", "/") + '"'
    else:
        here = os.path.abspath(__file__).replace("\\", "/")
        py = sys.executable.replace("\\", "/")
        cmd = f'"{py}" "{here}"'
    for key, value in (("diff.xlsx.textconv", f"{cmd} textconv"), ("diff.xlsx.binary", "true"),
                       ("diff.xlsx.command", f"{cmd} gitdiff"),
                       ("merge.xlsx.name", "xlgit cell-level merge"),
                       ("merge.xlsx.driver", f"{cmd} merge %O %A %B %P")):
        git("config", *where, key, value)


def install(scope="global"):
    root = repo_root()
    if scope == "repo" and not root:
        raise UserError("this folder isn't inside a git repository. Run it inside one, "
                        "or run `xlgit install` to set up every repository on this computer.")
    _set_drivers(["--global"] if scope == "global" else [])
    if scope == "global":
        _add_lines(global_attributes_file(), ATTR_LINES)
        print("xlgit is set up for every git repository on this computer.\n"
              "  git diff now lists changed cells in .xlsx/.xlsm files, and git merge combines\n"
              "  edits from different branches cell by cell.\n"
              "  Teammates run the same two commands once:  pip install xlgit  and  xlgit install")
    else:
        _add_lines(os.path.join(root, ".gitattributes"), ATTR_LINES)
        print("xlgit is set up for this repository. Commit .gitattributes so it applies to teammates too;\n"
              "  each of them runs `xlgit install` once on their own computer.")


def install_github():
    root = repo_root()
    if not root:
        raise UserError("run this inside the git repository you want pull request comments for.")
    path = os.path.join(root, ".github", "workflows", "excel-diff.yml")
    if os.path.exists(path):
        print(f"{os.path.relpath(path)} already exists; left as it is.")
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(WORKFLOW)
    _add_lines(os.path.join(root, ".gitattributes"), ATTR_LINES)
    rel = os.path.relpath(path, root).replace("\\", "/")
    print(f"Added {rel}. Commit and push it:\n"
          f"    git add {rel} .gitattributes\n"
          f"    git commit -m \"Show Excel changes on pull requests\"\n"
          f"    git push\n"
          f"  Every pull request that changes a workbook then gets a comment listing the changed cells.")


# ---------- demo ----------

DEMO_TRADES = [("Concrete", "m3", 120, 185), ("Framing", "m2", 2400, 6.5), ("Electrical", "m2", 2400, 4.25),
               ("Plumbing", "fixture", 18, 950), ("Finishes", "m2", 2400, 12)]


def _demo_book(path, trades):
    """A small cost estimate with a total and a chart, built the same way
    every time so the versions differ only where the story says."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate"
    ws.append(["Trade", "Unit", "Qty", "Rate", "Amount"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r, (trade, unit, qty, rate) in enumerate(trades, start=2):
        ws.append([trade, unit, qty, rate, f"=C{r}*D{r}"])
        ws[f"E{r}"].number_format = '"$"#,##0'
    total = len(trades) + 2
    ws[f"A{total}"] = "Total"
    ws[f"A{total}"].font = Font(bold=True)
    ws[f"E{total}"] = f"=SUM(E2:E{total - 1})"
    ws[f"E{total}"].number_format = '"$"#,##0'
    ws[f"E{total}"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 14
    chart = BarChart()
    chart.title = "Cost by trade"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=total - 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=total - 1))
    ws.add_chart(chart, "G2")
    wb.save(path)


def demo(folder=None, open_it=True):
    """Two estimators edit the same workbook on their own branches, and git
    merges them, in a throwaway repository. Nothing outside it is touched."""
    import shutil
    import tempfile
    if not shutil.which("git"):
        raise UserError("the demo needs git. Install it from https://git-scm.com and run `xlgit demo` again.")
    if folder:
        if os.path.exists(folder) and os.listdir(folder):
            raise UserError(f"{folder} isn't empty. Pick a new folder, or run `xlgit demo` without one.")
        os.makedirs(folder, exist_ok=True)
    root = os.path.abspath(folder or tempfile.mkdtemp(prefix="xlgit-demo-"))
    book = "estimate.xlsx"
    here = os.getcwd()
    os.chdir(root)
    try:
        git("init", "-q")
        git("symbolic-ref", "HEAD", "refs/heads/main")
        git("config", "user.name", "xlgit demo")
        git("config", "user.email", "demo@example.invalid")
        # The user's own git settings mustn't derail the story: fast-forward-only
        # merges, signed commits and hooks all belong to their real repositories.
        for key, value in (("merge.ff", "true"), ("commit.gpgsign", "false"),
                           ("core.hooksPath", os.path.join(root, ".git", "no-hooks"))):
            git("config", key, value)
        _set_drivers([])  # this repository only
        _add_lines(".gitattributes", ATTR_LINES)

        def commit(msg, trades):
            _demo_book(book, trades)
            git("add", "-A")
            git("commit", "-q", "-m", msg)

        def show(*args):
            out = git(*args).decode(errors="replace").rstrip()
            for line in out.splitlines():
                print("    " + line)

        print(f"A throwaway repository in {root}\n")
        print("1. Anna and Ben share estimate.xlsx: five trades, a total and a chart.")
        commit("Estimate for bid", DEMO_TRADES)

        roofing = ("Roofing", "m2", 1100, 38)
        anna = DEMO_TRADES[:2] + [roofing] + DEMO_TRADES[2:]
        git("checkout", "-q", "-b", "anna")
        commit("Add roofing", anna)
        print("\n2. Anna inserts a Roofing line on her branch. git diff main anna:")
        show("--no-pager", "diff", "main", "anna", "--", book)

        ben = [(t, u, 22 if t == "Plumbing" else q, 4.6 if t == "Electrical" else r) for t, u, q, r in DEMO_TRADES]
        git("checkout", "-q", "main")
        git("checkout", "-q", "-b", "ben")
        commit("New electrical rate, more fixtures", ben)
        print("\n3. Meanwhile Ben updates two numbers on his branch. git diff main ben:")
        show("--no-pager", "diff", "main", "ben", "--", book)

        print("\n4. Anna merges Ben's branch. Without xlgit this is a conflict on the whole file:")
        git("checkout", "-q", "anna")
        r = subprocess.run(["git", "merge", "--no-edit", "ben"], capture_output=True, text=True)
        for line in (r.stdout + r.stderr).splitlines():
            if line.strip():
                print("    " + line.strip())
        if r.returncode:
            raise UserError("the demo merge didn't go through cleanly, which it should. Please report it: " + ISSUES)

        cells = read_cells(book)["Estimate"]
        print("\n5. The merged estimate has both people's work. Ben's edits followed their rows down:")
        for row in range(2, len(anna) + 2):
            trade, qty, rate = cells.get(f"A{row}"), cells.get(f"C{row}"), cells.get(f"D{row}")
            note = {"Roofing": "   <- Anna", "Electrical": "   <- Ben's rate", "Plumbing": "   <- Ben's qty"}.get(trade, "")
            print(f"    row {row}  {trade:<11} qty {qty:>5}  rate {rate:>6}{note}")
        before = sum(q * r for _, _, q, r in DEMO_TRADES)
        after = sum(cells[f"C{row}"] * cells[f"D{row}"] for row in range(2, len(anna) + 2))
        print(f"    row {len(anna) + 2}  Total       {cells.get(f'E{len(anna) + 2}')}   "
              f"= ${after:,.0f} when Excel calculates it (was ${before:,.0f})")
        print("    The chart came through too. Open estimate.xlsx to see it.")

        print("\n6. The whole change as a grid, in your browser:")
        page = os.path.join(root, "changes.html")
        write_html([(book, Package(_git_blob("main", book)), Package.open(book))], page, open_it=open_it)
        print(f"\nPoke around in {root}: git log, xlgit diff, or open estimate.xlsx in Excel.\n"
              "To use xlgit on your own files, run `xlgit install` once.")
    finally:
        os.chdir(here)
    return root


def uninstall(scope="global"):
    where = ["--global"] if scope == "global" else []
    for section in ("diff.xlsx", "merge.xlsx"):
        git("config", *where, "--remove-section", section, check=False)
    if scope == "global":
        _remove_lines(global_attributes_file(), ATTR_LINES)
        print("xlgit is no longer set up globally. (Repositories set up with --repo keep their setup.)")
    else:
        root = repo_root()
        if root:
            _remove_lines(os.path.join(root, ".gitattributes"), ATTR_LINES)
        print("xlgit is no longer set up for this repository.")


def status_lines():
    driver = git("config", "--get", "merge.xlsx.driver", check=False).decode().strip()
    root = repo_root()
    out = []
    if not driver:
        out.append("[--] Not set up yet. Run:  xlgit install")
        return out
    exe = re.findall(r'"([^"]+)"', driver)
    if exe and not all(os.path.exists(p) for p in exe):
        out.append("[!!] Set up, but pointing at a Python or xlgit that no longer exists. Run:  xlgit install")
        return out
    scope = "every repository" if git("config", "--global", "--get", "merge.xlsx.driver",
                                          check=False).strip() else "this repository"
    out.append(f"[ok] Set up for {scope}.")
    if root:
        attr = git("check-attr", "merge", "--", "x.xlsx", check=False).decode()
        if "xlsx" not in attr:
            out.append("[!!] .xlsx files in this repository aren't using xlgit. Run:  xlgit install")
        if not os.path.exists(os.path.join(root, ".github", "workflows", "excel-diff.yml")):
            out.append("[  ] Pull request comments: not set up here. Optional:  xlgit install --github")
        else:
            out.append("[ok] Pull request comments are set up in this repository.")
    return out


HELP = f"""\
xlgit {__version__}: see and merge changes inside Excel files with git.

Try it (30 seconds, in a throwaway folder):
  xlgit demo                 two people edit one workbook, git merges it

Set up (once per computer):
  xlgit install              git diff and git merge understand .xlsx/.xlsm in every repository
  xlgit install --github     also comment the changed cells on this repository's pull requests

Everyday:
  git diff, git merge, git pull   work as usual, cell by cell
  xlgit diff                 what changed in your workbooks since the last commit
  xlgit diff FILE            ... in one workbook
  xlgit diff OLD NEW         compare any two workbooks
  xlgit diff --html          see the changes as a highlighted grid in your browser

Reporting a problem:
  xlgit scrub FILE...        a copy with every value made up, safe to attach to a bug report
  xlgit scrub --merge FILE   the three versions of a merge that went wrong, scrubbed, in one zip

More:
  xlgit install --repo       set up only the current repository
  xlgit uninstall [--repo]   undo the setup
  xlgit --version

Help and bug reports: {ISSUES}
"""


# ---------- diff against git ----------

def _git_blob(rev, path):
    r = subprocess.run(["git", "show", f"{rev}:{path}"], capture_output=True)
    return r.stdout if r.returncode == 0 else b""


def _changed_workbooks(*revs):
    names = git("diff", "--name-only", "-z", *revs).decode("utf-8", "replace").split("\0")
    return [n for n in names if n.lower().endswith((".xlsx", ".xlsm"))]


def _diff_blobs(old, new, title, markdown):
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        a, b = os.path.join(td, "old"), os.path.join(td, "new")
        with open(a, "wb") as f:
            f.write(old)
        with open(b, "wb") as f:
            f.write(new)
        if not markdown:
            print(f"=== {title} ===")
        try:
            return diff(a, b, markdown=markdown, title=title)
        except Exception as e:
            print(f"(couldn't compare {title}: {_reason(e)})\n")
            return 0


def worktree_changes(paths=()):
    """[(path, bytes at the last commit, bytes now)] for changed workbooks."""
    root = repo_root()
    if not root:
        raise UserError("not inside a git repository. To compare two files: xlgit diff OLD NEW")
    abspaths = [os.path.abspath(p) for p in paths]
    os.chdir(root)
    has_head = subprocess.run(["git", "rev-parse", "--verify", "-q", "HEAD"], capture_output=True).returncode == 0
    if paths:
        files = [os.path.relpath(p, root).replace("\\", "/") for p in abspaths]
    else:
        files = _changed_workbooks("HEAD") if has_head else []
        untracked = git("ls-files", "--others", "--exclude-standard", "-z").decode("utf-8", "replace").split("\0")
        files += [n for n in untracked if n.lower().endswith((".xlsx", ".xlsm")) and n not in files]
    return [(f, _git_blob("HEAD", f) if has_head else b"", open(f, "rb").read() if os.path.exists(f) else b"")
            for f in files]


def diff_worktree(paths=()):
    """Workbooks changed since the last commit (or new and untracked)."""
    changes = worktree_changes(paths)
    if not changes:
        print("No workbook changes since the last commit.")
        return 0
    changed = 0
    for f, old, new in changes:
        changed |= _diff_blobs(old, new, f, markdown=False)
    return changed


MARKER = "<!-- xlgit-pr-comment -->"


def pr_comment(base, head, limit=60000):
    """Markdown for a pull request: every workbook the PR changes."""
    mb = git("merge-base", base, head, check=False).decode().strip() or base
    out = io.StringIO()
    from contextlib import redirect_stdout
    files = _changed_workbooks(f"{mb}...{head}") if mb != base else _changed_workbooks(base, head)
    with redirect_stdout(out):
        print(MARKER)
        print("## Excel changes in this pull request\n")
        if not files:
            print("No workbook changes.")
        for f in files:
            _diff_blobs(_git_blob(mb, f), _git_blob(head, f), f, markdown=True)
    text = out.getvalue()
    if len(text) > limit:
        text = text[:limit].rsplit("\n", 1)[0] + "\n\n...cut short: GitHub comments have a size limit.\n"
    return text


# ---------- command line ----------

def main(argv):
    flags = {a for a in argv if a.startswith("--")}
    args = [a for a in argv if not a.startswith("--")]
    cmd = args.pop(0) if args else None
    if "--version" in flags or cmd == "version":
        print(f"xlgit {__version__}")
        return 0
    if cmd is None or cmd == "help" or "--help" in flags or "-h" in args:
        print(HELP)
        status = status_lines()
        for line in status:
            print(line)
        # Opened by double-clicking the standalone download: offer to set up.
        if cmd is None and getattr(sys, "frozen", False) and sys.stdin and sys.stdin.isatty():
            if status and not status[0].startswith("[ok]"):
                if input("\nSet up xlgit for git on this computer now? [Y/n] ").strip().lower() in ("", "y", "yes"):
                    install("global")
            input("\nPress Enter to close.")
        return 0
    if cmd == "textconv":
        # git diff and git log -p call this; never fail them over one bad file.
        try:
            textconv(args[0])
        except Exception as e:
            print(f"(xlgit couldn't read this file: {_reason(e)})")
        return 0
    if cmd == "gitdiff":
        # git diff runs this with: path old-file old-hex old-mode new-file new-hex new-mode
        if len(args) < 7:
            print(f"* Unmerged path {args[0] if args else ''}")
            return 0
        path, old, new = args[0], args[1], args[4]
        print(f"=== {path} ===")
        try:
            diff(old if old != "/dev/null" else "", new if new != "/dev/null" else "")
        except Exception as e:
            print(f"(xlgit couldn't compare this file: {_reason(e)})")
        sys.stdout.flush()
        return 0
    if cmd == "merge":
        if len(args) < 3:
            raise UserError("merge needs BASE OURS THEIRS (git passes these itself).")
        return merge(*args[:4])
    if cmd == "diff" and "--html" in flags:
        out = next((a.split("=", 1)[1] for a in flags if a.startswith("--out=")), None)
        if len(args) == 2:
            for p in args:
                if not os.path.exists(p):
                    raise UserError(f"no such file: {p}")
            items = [(os.path.basename(args[1]), Package.open(args[0]), Package.open(args[1]))]
        else:
            items = [(f, Package(old), Package(new)) for f, old, new in worktree_changes(args)]
            if not items:
                print("No workbook changes since the last commit.")
                return 0
        write_html(items, out, open_it="--no-open" not in flags)
        return 0
    if cmd == "diff":
        if len(args) == 2:
            title = next((a.split("=", 1)[1] for a in flags if a.startswith("--title=")), None)
            for p in args:
                if not os.path.exists(p):
                    raise UserError(f"no such file: {p}")
            rc = diff(args[0], args[1], markdown="--markdown" in flags, title=title)
            return 0 if "--markdown" in flags else rc
        return diff_worktree(args)
    if cmd == "pr-comment":
        if len(args) != 2:
            raise UserError("usage: xlgit pr-comment BASE HEAD")
        sys.stdout.write(pr_comment(*args))
        return 0
    if cmd == "install":
        if "--github" in flags:
            if not git("config", "--get", "merge.xlsx.driver", check=False).strip():
                install("repo" if "--repo" in flags else "global")
            install_github()
        else:
            install("repo" if "--repo" in flags else "global")
        return 0
    if cmd == "demo":
        demo(args[0] if args else None, open_it="--no-open" not in flags)
        return 0
    if cmd == "uninstall":
        uninstall("repo" if "--repo" in flags else "global")
        return 0
    if cmd == "scrub":
        if "--merge" in flags:
            if len(args) != 1:
                raise UserError("usage: xlgit scrub --merge FILE   (while a merge of FILE has a conflict)")
            out, sc = scrub_merge(args[0])
            print(f"Wrote {out} (base, yours and theirs, scrubbed).")
            print(_scrub_notice(sc))
            return 0
        if not args:
            raise UserError("usage: xlgit scrub FILE [FILE...]   or   xlgit scrub --merge FILE")
        for p in args:
            if not os.path.exists(p):
                raise UserError(f"no such file: {p}")
        written, sc = scrub_files(args)
        for w in written:
            print(f"Wrote {w}")
        print(_scrub_notice(sc))
        return 0
    if cmd == "status":
        for line in status_lines():
            print(line)
        return 0
    raise UserError(f"unknown command {cmd!r}. Run `xlgit` to see what it can do.")


def _utf8_output():
    """Cells can hold any language or symbol. Output going to git or a file
    is UTF-8, which git expects; a console that can't show a character (the
    Windows default) gets a stand-in instead of a crash."""
    for stream in (sys.stdout, sys.stderr):
        try:
            if stream.isatty():
                stream.reconfigure(errors="replace")
            else:
                stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError, OSError):
            pass


def cli():
    """Entry point for the installed `xlgit` command."""
    _utf8_output()
    try:
        code = main(sys.argv[1:])
    except UserError as e:
        say(f"xlgit: {e}")
        code = 2
    except KeyboardInterrupt:
        code = 130
    except Exception as e:
        if os.environ.get("XLGIT_DEBUG"):
            raise
        say(f"xlgit: {_reason(e)}\n  (set XLGIT_DEBUG=1 for details; please report it at {ISSUES})")
        code = 1
    sys.exit(code)


if __name__ == "__main__":
    cli()
